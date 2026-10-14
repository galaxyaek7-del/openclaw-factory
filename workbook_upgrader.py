"""Galaxy Forge -- turn header-only workbooks into working tools.

WHY THIS EXISTS
---------------
file_testing_engine.py (Cycle 87) opened every shipped .xlsx and found 19 of
77 contain ZERO formulas. Eleven share one generic skeleton (Deadlines /
Evidence log / Cost comparison / Correspondence log) whose only data row is
"Fill this row with your own real data". A buyer paying for a "tracker" got a
grid of column headers: no dropdowns, no date handling, no totals, nothing
that tells them what is overdue. That is the gap between a file and a tool.

WHAT IT DOES (deterministic, no AI call, no invented content)
--------------------------------------------------------------
Every upgrade is derived from the column headers the product already has. No
example data, numbers or claims are added -- only structure that computes on
the buyer's own entries:

  * classifies each column (date / due date / money / yes-no / status /
    outcome / rating / number / text) from its header
  * dropdowns for yes-no, status and outcome columns (non-blocking: a buyer
    may still type their own value)
  * date and currency number formats on 200 entry rows
  * a "Days left" column beside every due-date column, red when overdue and
    amber within 7 days, blank once the row's status is Done
  * header-declared formulas made live (e.g. "Tonnes/yr (=C*D/1e6)")
  * a Summary sheet: entries per sheet, money totals, status counts, overdue
    and due-this-week counts, yes-rates -- all live formulas
  * a How to use sheet; the "fill this row" placeholder moves there
  * header styling, frozen header row, auto-filter, column widths

It refuses to touch a workbook it has already upgraded (marker in the file's
keywords), and treats README / Instructions sheets as text.

LIMITATION, STATED
------------------
No spreadsheet recalculation engine (LibreOffice, formulas, pycel) is
installed here, so formulas are verified structurally (syntax, sheet and
column references) by the tests, not by computing them.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

MARKER = "gf-workbook-upgrade-v1"
ENTRY_ROWS = 200
TEXT_SHEET_NAMES = {"readme", "instructions", "how to use", "notes", "summary", "dashboard"}

YES_NO = ["Yes", "No", "N/A"]
Y_N = ["Y", "N"]
PASS_FAIL = ["Pass", "Fail", "Partial", "N/A"]
STATUS = ["Not started", "In progress", "Waiting on other party", "Blocked", "Done"]
OUTCOME = ["Pending", "Won", "Lost", "Partially won", "Withdrawn"]
DONE_VALUES = ("Done",)

HEADER_FILL = "1F3864"
INPUT_HINT_FILL = "F2F2F2"
COMPUTED_FILL = "E8EEF7"

_PLACEHOLDER = re.compile(r"fill this row|nothing here is pre-?filled", re.I)
_HEADER_FORMULA = re.compile(r"\(=\s*([^)]+)\)")


# --- column classification -------------------------------------------------
@dataclass
class Column:
    index: int          # 1-based
    letter: str
    header: str
    kind: str           # date|due|money|yesno|status|outcome|rating|number|formula|text
    options: list = field(default_factory=list)
    formula: str = ""


_ANNOTATION = re.compile(r"\s*\((illustrative|example|optional|d-number)\)", re.I)


def classify(header: str) -> tuple[str, list]:
    # "Rate (illustrative)" is a money column; the note says the sample value
    # is made up, not that the column is text. Typed suffixes such as "(Y/N)",
    # "($)", "(1-5)", "(hours)" and "(=C*D)" are kept and still classify.
    h = _ANNOTATION.sub("", (header or "")).strip().lower()
    if not h:
        return "text", []
    if _HEADER_FORMULA.search(header):
        return "formula", []
    if "(y/n)" in h:
        return "yesno", Y_N
    if h in ("pass?", "pass/fail") or h.startswith("pass?"):
        return "yesno", PASS_FAIL
    if any(k in h for k in ("deadline", "due", "next re-audit", "renewal",
                            "review date", "remediation date", "expires")):
        return "due", []
    # word boundary: "AR mandated?" and "Statement updated?" are not dates
    if re.search(r"\bdates?\b", h) or h == "in place since":
        return "date", []
    if re.search(r"stars?\b.*\(1-5\)|^stars?$", h):
        return "rating", []
    if "%" in h or "percent" in h:
        return "number", []
    if any(k in h for k in ("($)", "amount", "cost", "fee", "price", "value",
                            "recovered", "budget", "position ($)", "billed")) \
            or h in ("rate", "paid") or re.search(r"\brate\b", h):
        return "money", []
    if h.endswith("?") or any(k in h for k in ("eligible", "attached", "obtained",
                                               "disclosed", "retention set",
                                               "reviewed")):
        return "yesno", YES_NO
    if any(k in h for k in ("outcome", "win/loss", "result")):
        return "outcome", OUTCOME
    if h.startswith("status") or h == "stage":
        return "status", STATUS
    if any(k in h for k in ("(hours)", "(min)", "(g)", "days to", "units",
                            "missed calls", "jobs won", "jobs lost", "rounds",
                            "round used", "new fails", "minutes", "(days")):
        return "number", []
    return "text", []


def _q(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _header_formula(expr: str, row: int) -> str:
    # "C*D/1e6" -> "C5*D5/1000000". Scientific notation is expanded FIRST, so
    # the "e6" in "1e6" can never be mistaken for a cell reference; refs are
    # collected from the substitution itself, not re-scanned from the output.
    expr = re.sub(r"\d+(?:\.\d+)?[eE][+-]?\d+",
                  lambda m: format(float(m.group(0)), "f").rstrip("0").rstrip("."),
                  expr.strip())
    refs: list[str] = []

    def _ref(m):
        refs.append(f"{m.group(1)}{row}")
        return refs[-1]

    body = re.sub(r"(?<![A-Za-z0-9.])([A-Z])(?![A-Za-z0-9])", _ref, expr)
    blank = ",".join(f'{r}=""' for r in refs) or 'FALSE'
    return f'=IF(OR({blank}),"",IFERROR({body},""))'


# --- upgrade ---------------------------------------------------------------
@dataclass
class UpgradeReport:
    path: str
    skipped: str = ""
    sheets: dict = field(default_factory=dict)
    placeholders_moved: int = 0
    formulas_added: int = 0
    validations_added: int = 0


def _is_text_sheet(ws) -> bool:
    if ws.title.strip().lower() in TEXT_SHEET_NAMES:
        return True
    return ws.max_column <= 1 and ws.max_row > 1 and ws.title.lower().startswith("read")


def _headers(ws) -> list[str]:
    return [("" if c.value is None else str(c.value)) for c in ws[1]]


def upgrade_workbook(src, dst=None) -> UpgradeReport:
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    src = Path(src)
    dst = Path(dst) if dst else src
    rep = UpgradeReport(str(src))
    wb = openpyxl.load_workbook(str(src))
    if MARKER in (wb.properties.keywords or ""):
        rep.skipped = "already upgraded"
        return rep

    hfill = PatternFill("solid", fgColor=HEADER_FILL)
    cfill = PatternFill("solid", fgColor=COMPUTED_FILL)
    hfont = Font(bold=True, color="FFFFFF")
    red = PatternFill("solid", fgColor="F8CBAD")
    amber = PatternFill("solid", fgColor="FFE699")

    placeholders: list[str] = []
    plan: dict[str, list[Column]] = {}
    last_row: dict[str, int] = {}

    for ws in list(wb.worksheets):
        if _is_text_sheet(ws):
            continue
        heads = _headers(ws)
        if not any(heads):
            continue
        # move "fill this row" placeholder rows out of the data area
        for r in range(ws.max_row, 1, -1):
            vals = [c.value for c in ws[r]]
            nonempty = [v for v in vals if v not in (None, "")]
            if len(nonempty) == 1 and isinstance(nonempty[0], str) \
                    and _PLACEHOLDER.search(nonempty[0]):
                placeholders.append(nonempty[0])
                ws.delete_rows(r)
        while heads and not heads[-1]:
            heads.pop()
        cols = []
        for i, h in enumerate(heads, 1):
            kind, opts = classify(h)
            col = Column(i, get_column_letter(i), h, kind, opts)
            if kind == "formula":
                col.formula = _HEADER_FORMULA.search(h).group(1)
            cols.append(col)
        status_col = next((c for c in cols if c.kind == "status"), None)

        first = 2
        last = max(ws.max_row, 1) + ENTRY_ROWS
        last_row[ws.title] = last
        added = []

        # computed "Days left" beside each due column
        next_idx = len(cols) + 1
        for c in [c for c in cols if c.kind == "due"]:
            letter = get_column_letter(next_idx)
            dl = Column(next_idx, letter, f"Days left ({c.header})"[:60], "computed")
            next_idx += 1
            ws.cell(row=1, column=dl.index, value=dl.header)
            for r in range(first, last + 1):
                due = f"{c.letter}{r}"
                done = (f',{status_col.letter}{r}="Done"' if status_col else "")
                ws.cell(row=r, column=dl.index,
                        value=f'=IF(OR({due}=""{done}),"",IFERROR(INT({due}-TODAY()),""))').fill = cfill
                rep.formulas_added += 1
            rng = f"{letter}{first}:{letter}{last}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({letter}{first}),{letter}{first}<0)'], fill=red))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(ISNUMBER({letter}{first}),{letter}{first}>=0,{letter}{first}<=7)'], fill=amber))
            added.append(dl)

        for c in cols:
            rng = f"{c.letter}{first}:{c.letter}{last}"
            if c.options:
                dv = DataValidation(type="list", formula1='"' + ",".join(c.options) + '"',
                                    allow_blank=True, showErrorMessage=False)
                dv.promptTitle = c.header[:32]
                dv.prompt = "Pick from the list or type your own value."
                dv.showInputMessage = True
                ws.add_data_validation(dv)
                dv.add(rng)
                rep.validations_added += 1
            if c.kind == "rating":
                dv = DataValidation(type="whole", operator="between", formula1="1",
                                    formula2="5", allow_blank=True)
                dv.error = "Enter a whole number from 1 to 5."
                ws.add_data_validation(dv)
                dv.add(rng)
                rep.validations_added += 1
            for r in range(first, last + 1):
                cell = ws.cell(row=r, column=c.index)
                if c.kind in ("date", "due"):
                    cell.number_format = "yyyy-mm-dd"
                elif c.kind == "money":
                    cell.number_format = '#,##0.00'
                elif c.kind == "formula":
                    if cell.value in (None, ""):
                        cell.value = _header_formula(c.formula, r)
                        cell.fill = cfill
                        cell.number_format = "#,##0.000"
                        rep.formulas_added += 1

        allcols = cols + added
        for c in allcols:
            cell = ws.cell(row=1, column=c.index)
            cell.font = hfont
            cell.fill = hfill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[c.letter].width = max(14, min(40, len(c.header) + 6))
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(allcols))}{last}"
        plan[ws.title] = allcols
        rep.sheets[ws.title] = {c.header: c.kind for c in allcols}

    # --- Summary sheet -----------------------------------------------------
    # a builder-made "Summary" (cohort_workbook.py) is kept; ours becomes "Dashboard"
    sm = wb.create_sheet("Dashboard" if "Summary" in wb.sheetnames else "Summary", 0)
    sm["A1"] = "Summary -- updates automatically from your entries"
    sm["A1"].font = Font(bold=True, size=14)
    sm.column_dimensions["A"].width = 46
    sm.column_dimensions["B"].width = 18
    row = 3
    for title, cols in plan.items():
        q, last = _q(title), last_row[title]
        sm.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
        row += 1
        key = cols[0]
        sm.cell(row=row, column=1, value="Rows in use")
        sm.cell(row=row, column=2, value=f"=COUNTA({q}!{key.letter}2:{key.letter}{last})")
        row += 1
        status = next((c for c in cols if c.kind == "status"), None)
        for c in cols:
            ref = f"{q}!{c.letter}2:{c.letter}{last}"
            if c.kind == "money":
                sm.cell(row=row, column=1, value=f"Total: {c.header}")
                cell = sm.cell(row=row, column=2, value=f"=SUM({ref})")
                cell.number_format = "#,##0.00"
                row += 1
            elif c.kind == "due":
                open_cond = (f',{q}!{status.letter}2:{status.letter}{last},"<>Done"'
                             if status else "")
                sm.cell(row=row, column=1, value=f"Overdue: {c.header}")
                sm.cell(row=row, column=2,
                        value=f'=COUNTIFS({ref},"<"&TODAY(){open_cond})')
                row += 1
                sm.cell(row=row, column=1, value=f"Due in next 7 days: {c.header}")
                sm.cell(row=row, column=2,
                        value=f'=COUNTIFS({ref},">="&TODAY(),{ref},"<="&TODAY()+7{open_cond})')
                row += 1
            elif c.kind in ("status", "outcome"):
                for opt in c.options:
                    sm.cell(row=row, column=1, value=f"{c.header}: {opt}")
                    sm.cell(row=row, column=2, value=f'=COUNTIF({ref},"{opt}")')
                    row += 1
            elif c.kind == "yesno":
                yes = c.options[0]
                sm.cell(row=row, column=1, value=f"{c.header}: share '{yes}'")
                cell = sm.cell(row=row, column=2,
                               value=f'=IFERROR(COUNTIF({ref},"{yes}")/COUNTA({ref}),"")')
                cell.number_format = "0%"
                row += 1
            elif c.kind == "rating":
                sm.cell(row=row, column=1, value=f"Average: {c.header}")
                cell = sm.cell(row=row, column=2, value=f'=IFERROR(AVERAGE({ref}),"")')
                cell.number_format = "0.0"
                row += 1
        row += 1
    for r in range(3, row):
        if isinstance(sm.cell(row=r, column=2).value, str):
            rep.formulas_added += 1

    # --- How to use sheet --------------------------------------------------
    hu = wb.create_sheet("How to use", 0)
    hu.column_dimensions["A"].width = 100
    lines = [
        ("How to use this workbook", True),
        ("", False),
        ("1. Enter one item per row on each tracker sheet. Rows 2-%d are ready for entries." % (ENTRY_ROWS + 1), False),
        ("2. Columns with a dropdown offer standard values; you may still type your own.", False),
        ("3. Enter dates as dates (e.g. 2026-11-30) so deadlines can be calculated.", False),
        ("4. Shaded 'Days left' columns calculate themselves: red = overdue, amber = due within 7 days.", False),
        ("   They go blank once you set the row's Status to Done.", False),
        ("5. The Summary sheet totals amounts, counts statuses and flags what is overdue -- live.", False),
        ("6. Use the filter arrows in each header row to sort or narrow your list.", False),
        ("", False),
        ("Nothing in this workbook is pre-filled with example figures. Every number comes from your own entries.", False),
    ]
    for p in dict.fromkeys(placeholders):
        lines.append((p, False))
    for i, (text, bold) in enumerate(lines, 1):
        c = hu.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if bold else 11)
    rep.placeholders_moved = len(placeholders)

    kw = (wb.properties.keywords or "").strip()
    wb.properties.keywords = (kw + " " + MARKER).strip()
    wb.active = 0
    wb.save(str(dst))
    return rep


if __name__ == "__main__":  # pragma: no cover
    import json
    import sys
    for arg in sys.argv[1:]:
        r = upgrade_workbook(arg)
        print(json.dumps(r.__dict__, indent=1, default=str))
