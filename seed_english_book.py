"""OpenClaw Factory — English Book Seed

Standalone. Produces ONE English DOCX manuscript ready for KDP.
Does NOT touch the live factory pipeline. Safe to run/delete.

Usage:
  python seed_english_book.py --niche "Morning Focus Journal for Remote Workers" --price 6.99

Output:
  seeds/<slug>.docx  — the manuscript
  seeds/<slug>.json  — metadata (niche, title, price, chapters, wordcount)
"""

import argparse
import json
import os
import re
import sys
import time
import urllib.request
import urllib.error
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt

# Reuse the real parser (Task 15-A found it already language-tolerant — its
# header matching checks English keywords too, e.g. 'chapter'/'conclusion'/
# 'intro' in header.lower()) rather than duplicate it. Importing book_generator
# is safe in this environment (reportlab is installed, confirmed before
# writing this file) — its only side effect at import time is defining
# functions/constants, no PDF is drawn until a function is actually called.
# Per the task's own constraint, ONLY _parse_sectioned_book is used from it —
# no other book_generator helper (e.g. get_groq_key) is imported, so the
# small Groq-call/key-reading logic below is its own standalone copy, not a
# reuse of book_generator's version.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from book_generator import _parse_sectioned_book  # noqa: E402


GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"
FACTORY_DIR = os.path.dirname(os.path.abspath(__file__))
SEEDS_DIR = os.path.join(FACTORY_DIR, "seeds")


def _get_groq_key():
    key = os.environ.get("GROQ_KEY")
    if key:
        return key
    env_path = os.path.join(FACTORY_DIR, ".env")
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line.startswith("GROQ_KEY"):
                    return line.split("=", 1)[1].strip()
    return None


def build_prompt(niche, target_chapters=8, target_words_per_chapter=1200):
    system = (
        "You are a professional non-fiction ghostwriter producing publish-ready "
        "manuscripts for Amazon KDP. Write in English (US). Practical, no fluff, "
        "no filler sentences, no generic platitudes. Voice: second-person "
        "('you'), short paragraphs (2-4 sentences), plain confident language a "
        "busy reader can act on immediately.\n\n"
        "Hard content rules — refuse silently and just avoid these, do not "
        "explain the refusal in the output:\n"
        "- Never mention or reference real trademarked brands, characters, "
        "franchises, or companies (e.g. Disney, Marvel, Pokemon, Nintendo, "
        "Harry Potter, Nike, or any other real trademark).\n"
        "- Never give financial investment advice, trading strategies, stock "
        "picks, or any 'get rich quick' / 'guaranteed returns' framing.\n"
        "- Never give medical advice, diagnoses, cures, or treatment claims of "
        "any kind, even framed as general wellness tips.\n"
        "- No scam-adjacent language (nothing implying guaranteed/secret/"
        "insider results).\n"
        "- No sexual or adult content."
    )

    chapter_markers = "\n".join(
        f"# CHAPTER {i}: <a specific, concrete chapter title>\n"
        f"(~{target_words_per_chapter} words. Include ONE concrete, specific "
        f"example and ONE clear action step the reader can do today.)"
        for i in range(1, target_chapters + 1)
    )

    user = f"""Write a complete, publish-ready non-fiction book manuscript.

NICHE / TOPIC: {niche}
NUMBER OF CHAPTERS: {target_chapters}
TARGET LENGTH PER CHAPTER: ~{target_words_per_chapter} words

Output ONLY the manuscript, in EXACTLY this literal Markdown-heading format, \
in this exact order, with a leading "#" on every section heading line \
(this is a hard machine-parsed format — do not deviate):

# INTRODUCTION
(150-250 words. Hook the reader, state what they'll be able to do after \
reading this book.)
{chapter_markers}
# CONCLUSION
(100-150 words. Summarize the transformation, end with encouragement to act.)

Do not add any text before "# INTRODUCTION" or after the conclusion body. Do \
not add numbering, explanations, or commentary outside these sections. Every \
chapter must contain real, specific, useful content — never a placeholder or \
a restatement of the chapter title."""

    return system, user


def call_groq(prompt, model="llama-3.1-8b-instant", max_tokens=4096, timeout=60):
    system, user = prompt
    key = _get_groq_key()
    if not key:
        raise RuntimeError("GROQ_KEY not found in environment or .env — cannot call Groq")

    payload = json.dumps({
        "model": model,
        "max_tokens": max_tokens,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    }).encode("utf-8")
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {key}",
        "User-Agent": "Mozilla/5.0 (OpenClaw-Factory-EnglishSeed)",
        "Accept": "application/json",
    }

    last_error = None
    for attempt in (1, 2):  # one retry on network error, per spec
        try:
            req = urllib.request.Request(GROQ_API_URL, data=payload, method="POST", headers=headers)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                result = json.loads(r.read().decode("utf-8"))
            text = result["choices"][0]["message"]["content"]
            tokens = None
            usage = result.get("usage") or {}
            if isinstance(usage, dict) and "total_tokens" in usage:
                tokens = usage["total_tokens"]
            return text, tokens
        except urllib.error.HTTPError as e:
            # Auth/bad-request errors are not transient — retrying can't fix a bad key.
            last_error = e
            if e.code in (400, 401, 403):
                break
        except Exception as e:
            last_error = e
        if attempt == 1:
            time.sleep(1)

    raise RuntimeError(f"Groq call failed after retry: {last_error}")


def parse_book(raw_text, niche, expected_chapters):
    parsed = _parse_sectioned_book(raw_text, expected_chapters)
    return {
        "title": niche.strip(),
        "introduction": parsed.get("introduction", ""),
        "chapters": [
            {"title": ch.get("title", ""), "body": ch.get("content", "")}
            for ch in parsed.get("chapters", [])
        ],
        "conclusion": parsed.get("conclusion", ""),
    }


def _set_run_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold


def _add_body_paragraphs(doc, text):
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(para)
        _set_run_font(run, "Calibri", 11)


def write_docx(book_dict, out_path):
    doc = Document()

    h1 = doc.add_heading(level=1)
    _set_run_font(h1.add_run(book_dict["title"]), "Calibri", 16, bold=True)

    doc.add_paragraph()

    if book_dict.get("introduction"):
        h2 = doc.add_heading(level=2)
        _set_run_font(h2.add_run("Introduction"), "Calibri", 14, bold=True)
        _add_body_paragraphs(doc, book_dict["introduction"])

    for i, ch in enumerate(book_dict.get("chapters", []), 1):
        h2 = doc.add_heading(level=2)
        title = ch.get("title") or f"Chapter {i}"
        # Chapter title from the parser already includes a "CHAPTER N:" prefix
        # (it's the full matched header text) — avoid double-prefixing it.
        label = title if re.match(r'(?i)^chapter\s*\d+', title) else f"Chapter {i}: {title}"
        _set_run_font(h2.add_run(label), "Calibri", 14, bold=True)
        _add_body_paragraphs(doc, ch.get("body", ""))

    if book_dict.get("conclusion"):
        h2 = doc.add_heading(level=2)
        _set_run_font(h2.add_run("Conclusion"), "Calibri", 14, bold=True)
        _add_body_paragraphs(doc, book_dict["conclusion"])

    doc.save(out_path)


def _slugify(text):
    text = text.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-")
    return text or "untitled"


def _wordcount(book_dict):
    parts = [book_dict.get("introduction", ""), book_dict.get("conclusion", "")]
    parts += [ch.get("body", "") for ch in book_dict.get("chapters", [])]
    return sum(len(p.split()) for p in parts if p)


_ARABIC_RE = re.compile(r"[؀-ۿ]")


def _sanity_check(book_dict):
    issues = []
    if not book_dict.get("chapters"):
        issues.append("no chapters parsed")
    full_text = " ".join(
        [book_dict.get("title", ""), book_dict.get("introduction", ""), book_dict.get("conclusion", "")]
        + [c.get("title", "") + " " + c.get("body", "") for c in book_dict.get("chapters", [])]
    )
    arabic_hits = _ARABIC_RE.findall(full_text)
    if arabic_hits:
        issues.append(f"{len(arabic_hits)} Arabic character(s) leaked into the manuscript")
    return issues


def main():
    parser = argparse.ArgumentParser(description="Generate one English DOCX book seed for KDP.")
    parser.add_argument("--niche", required=True, help="Book niche/topic (also used as the title).")
    parser.add_argument("--price", type=float, default=6.99)
    parser.add_argument("--chapters", type=int, default=8)
    parser.add_argument("--force", action="store_true", help="Overwrite existing output files.")
    args = parser.parse_args()

    stage = "setup"
    try:
        os.makedirs(SEEDS_DIR, exist_ok=True)
        slug = _slugify(args.niche)
        docx_path = os.path.join(SEEDS_DIR, f"{slug}.docx")
        json_path = os.path.join(SEEDS_DIR, f"{slug}.json")

        if not args.force:
            existing = [p for p in (docx_path, json_path) if os.path.exists(p)]
            if existing:
                print(f"ERROR at stage 'setup': refusing to overwrite existing file(s): {existing}. Use --force.")
                sys.exit(1)

        stage = "build_prompt"
        prompt = build_prompt(args.niche, target_chapters=args.chapters)

        stage = "call_groq"
        model = "llama-3.1-8b-instant"
        raw_text, tokens_used = call_groq(prompt, model=model)

        stage = "parse_book"
        book = parse_book(raw_text, args.niche, args.chapters)

        stage = "write_docx"
        write_docx(book, docx_path)

        stage = "sanity_check"
        issues = _sanity_check(book)
        wordcount = _wordcount(book)

        stage = "write_metadata"
        metadata = {
            "niche": args.niche,
            "title": book["title"],
            "price_usd": args.price,
            "platform": "kdp_ebook",
            "language": "en",
            "chapters": len(book.get("chapters", [])),
            "wordcount": wordcount,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "model": model,
            "groq_tokens_used": tokens_used,
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(metadata, f, ensure_ascii=False, indent=2)

        verdict = "OK — no issues found" if not issues else "ISSUES: " + "; ".join(issues)

        print(f"docx: {docx_path}")
        print(f"json: {json_path}")
        print(f"wordcount: {wordcount}")
        print(f"chapters parsed: {len(book.get('chapters', []))}")
        print(f"verdict: {verdict}")

        if issues:
            sys.exit(1)

    except Exception as e:
        print(f"FAILED at stage '{stage}': {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
