"""Tests for workbook_upgrader.py and pdf_glyph_repair.py.

Both modules rewrite files customers have paid for, so each test pins a way a
repair could quietly make a product worse: misclassifying a column, mistaking
"1e6" for a cell, touching a real bullet, or writing a PDF whose text changed.
"""
import re
import sys
import tempfile
import importlib.util
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import file_testing_engine as fte
import pdf_glyph_repair as pgr
import workbook_upgrader as wu

TMP = Path(tempfile.gettempdir()) / "gf_quality_repairs"


def _p(name):
    TMP.mkdir(parents=True, exist_ok=True)
    return TMP / name


def _skeleton(path, sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    wb.save(path)


PLACEHOLDER = "Fill this row with your own real data — nothing here is pre-filled."


class TestClassify(unittest.TestCase):
    def test_words_containing_date_are_not_dates(self):
        # real headers from shipped workbooks that a substring match got wrong
        self.assertEqual(wu.classify("AR mandated?")[0], "yesno")
        self.assertEqual(wu.classify("Statement updated?")[0], "yesno")

    def test_real_headers(self):
        cases = {
            "Response due": "due", "Next re-audit": "due", "Due date": "due",
            "Date": "date", "Done + retest date": "date", "In place since": "date",
            "Dispute amount": "money", "Cost ($)": "money", "Fee": "money",
            "Paid (Y/N)": "yesno", "Evidence attached": "yesno",
            "Status": "status", "Win/loss": "outcome", "Outcome": "outcome",
            "Stars (1-5)": "rating", "Reply time (hours)": "number",
            "Tonnes/yr (=C*D/1e6)": "formula", "Notes": "text", "Owner": "text",
            "Decision": "text",  # build-vs-buy, not won/lost
        }
        for h, want in cases.items():
            self.assertEqual(wu.classify(h)[0], want, h)

    def test_annotation_suffix_does_not_hide_the_type(self):
        # real headers from the live freelancer workbook
        cases = {
            "Rate (illustrative)": "money", "Overage rate per round (illustrative)": "money",
            "Extra billed": "money", "Minutes spent": "number",
            "Schedule impact (days, illustrative)": "number",
            "Rounds included (illustrative)": "number", "Request ID (illustrative)": "text",
            "Interest rate (%)": "number",          # a percent is not money
            "Paid (Y/N)": "yesno", "Cost ($)": "money",  # typed suffixes still win
        }
        for h, want in cases.items():
            self.assertEqual(wu.classify(h)[0], want, h)


class TestHeaderFormula(unittest.TestCase):
    def test_scientific_notation_is_not_a_cell_reference(self):
        f = wu._header_formula("C*D/1e6", 5)
        self.assertEqual(f, '=IF(OR(C5="",D5=""),"",IFERROR(C5*D5/1000000,""))')
        self.assertNotIn("E6", f)


def _formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate, c.value


class TestUpgradeWorkbook(unittest.TestCase):
    def setUp(self):
        self.path = _p("generic_tracker.xlsx")
        _skeleton(self.path, [
            ("Deadlines", ["Dispute ID", "Response due", "Owner", "Status", "Notes"], [[PLACEHOLDER]]),
            ("Evidence log", ["Dispute ID", "Evidence attached", "Outcome"], [[PLACEHOLDER]]),
            ("Cost comparison", ["Dispute amount", "Fee", "Win/loss"], [[PLACEHOLDER]]),
        ])
        self.before = fte.test_file(self.path)

    def test_header_only_workbook_fails_before_and_passes_after(self):
        self.assertEqual(self.before.status, fte.FAIL)
        wu.upgrade_workbook(self.path)
        after = fte.test_file(self.path)
        self.assertEqual(after.status, fte.PASS, after.problems)

    def test_structure(self):
        import openpyxl
        rep = wu.upgrade_workbook(self.path)
        self.assertEqual(rep.placeholders_moved, 3)
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb.sheetnames[:2], ["How to use", "Summary"])
        dl = wb["Deadlines"]
        # placeholder left the data area; moved to How to use
        self.assertIsNone(dl["A2"].value)
        self.assertIn(PLACEHOLDER, [c.value for c in wb["How to use"]["A"]])
        # Days left beside the due column, blank once Done
        self.assertTrue(dl["F1"].value.startswith("Days left"))
        self.assertEqual(dl["F2"].value, '=IF(OR(B2="",D2="Done"),"",IFERROR(INT(B2-TODAY()),""))')
        self.assertEqual(dl["B2"].number_format, "yyyy-mm-dd")
        # dropdowns are non-blocking suggestions
        dvs = dl.data_validations.dataValidation
        self.assertTrue(any("Done" in (d.formula1 or "") for d in dvs))
        self.assertTrue(all(d.showErrorMessage is False for d in dvs if d.type == "list"))
        summary = [c.value for c in wb["Summary"]["B"] if c.value]
        self.assertIn("=SUM('Cost comparison'!A2:A201)", summary)
        self.assertTrue(any("COUNTIFS('Deadlines'!B2:B201,\"<\"&TODAY()" in s for s in summary))

    def test_every_formula_is_structurally_sound(self):
        import openpyxl
        wu.upgrade_workbook(self.path)
        wb = openpyxl.load_workbook(self.path)
        names = set(wb.sheetnames)
        for sheet, coord, f in _formulas(wb):
            self.assertEqual(f.count("("), f.count(")"), f"{sheet}!{coord}: {f}")
            self.assertEqual(f.count('"') % 2, 0, f"{sheet}!{coord}: {f}")
            for ref in re.findall(r"'((?:[^']|'')+)'!", f):
                self.assertIn(ref.replace("''", "'"), names, f)
            self.assertNotRegex(f, r"#(REF!|NAME\?)")

    def test_idempotent(self):
        wu.upgrade_workbook(self.path)
        self.assertEqual(wu.upgrade_workbook(self.path).skipped, "already upgraded")

    def test_prefilled_checklist_rows_are_kept(self):
        p = _p("checklist.xlsx")
        _skeleton(p, [("Eligibility", ["Check", "Your answer", "Pass?"],
                       [["Not at fault", None, None], ["Within 3 years", None, None]])])
        wu.upgrade_workbook(p)
        import openpyxl
        ws = openpyxl.load_workbook(p)["Eligibility"]
        self.assertEqual([ws["A2"].value, ws["A3"].value], ["Not at fault", "Within 3 years"])

    def test_readme_sheet_untouched(self):
        p = _p("with_readme.xlsx")
        _skeleton(p, [("Log", ["Date", "Status"], []),
                      ("README", ["Ask every customer within 24 hours."], [["Reply to every review."]])])
        wu.upgrade_workbook(p)
        import openpyxl
        ws = openpyxl.load_workbook(p)["README"]
        self.assertEqual(ws.max_column, 1)
        self.assertEqual(ws["A2"].value, "Reply to every review.")


def _box_pdf(path, body):
    """Build a PDF the way the factory did before text_sanitizer existed."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    st = getSampleStyleSheet()
    SimpleDocTemplate(str(path), pagesize=letter).build(
        [Paragraph(t, st["BodyText"]) for t in body])


class TestPdfGlyphRepair(unittest.TestCase):
    def test_in_word_boxes_repaired_and_verified(self):
        p = _p("boxes.pdf")
        _box_pdf(p, ["Post‐Removal plan within the 30‑day window.",
                     "A well‐structured, third‐party verified appeal."])
        self.assertEqual(pgr.count_in_word_boxes(p), 4)
        self.assertEqual(fte.test_pdf(p).detail["in_word_box_glyphs"], 4)
        rep = pgr.repair_pdf(p)
        self.assertTrue(rep.written, rep.reason)
        self.assertEqual((rep.boxes_before, rep.runs_rewritten, rep.boxes_after), (4, 4, 0))
        self.assertEqual(pgr.count_in_word_boxes(p), 0)
        from pypdf import PdfReader
        text = PdfReader(str(p)).pages[0].extract_text()
        self.assertIn("Post-Removal", text)
        self.assertIn("30-day", text)

    def test_boxes_one_character_apart(self):
        # found in 3 shipped PDFs: "4-7-8" and "as-a-service" style runs
        p = _p("adjacent.pdf")
        _box_pdf(p, ["Breathe 4‐7‐8 and use function‐as‐a‐service."])
        rep = pgr.repair_pdf(p)
        self.assertTrue(rep.written, rep.reason)
        self.assertEqual(rep.boxes_after, 0)
        from pypdf import PdfReader
        self.assertIn("as-a-service", PdfReader(str(p)).pages[0].extract_text())

    def test_real_bullet_is_never_touched(self):
        p = _p("bullet.pdf")
        _box_pdf(p, ["■ Keep receipts", "Done ■ next"])
        rep = pgr.repair_pdf(p)
        self.assertFalse(rep.written)
        self.assertEqual(rep.reason, "no in-word boxes")

    def test_clean_pdf_untouched(self):
        p = _p("clean.pdf")
        _box_pdf(p, ["A plain-hyphen sentence."])
        before = p.read_bytes()
        self.assertFalse(pgr.repair_pdf(p).written)
        self.assertEqual(p.read_bytes(), before)

    def test_file_tester_fails_a_boxed_pdf(self):
        p = _p("boxed_for_tester.pdf")
        _box_pdf(p, ["Post‐Removal"])
        r = fte.test_pdf(p)
        self.assertEqual(r.status, fte.FAIL)
        self.assertTrue(any("black-box" in x for x in r.problems))


class TestBlockingProblems(unittest.TestCase):
    def test_layout_heuristics_are_advisory(self):
        r = fte.CheckResult("x.pdf", "pdf", fte.FAIL, {},
                            ["layout MARGIN_BREACH p1: text at (0, 18)"])
        self.assertEqual(fte.blocking_problems(r), [])

    def test_unknown_never_blocks(self):
        r = fte.CheckResult("x.pdf", "pdf", fte.UNKNOWN, {}, ["does not parse"])
        self.assertEqual(fte.blocking_problems(r), [])

    def test_title_overlapping_subtitle_is_detected(self):
        # The real defect: page 1 title wraps onto its subtitle. The live file
        # was repaired 2026-10-17 (scripts/fix_pdf_title_overlap.py); the
        # original, kept for exactly this purpose, is the defect fixture now,
        # and the repaired file must come back clean.
        books = Path(__file__).resolve().parent.parent / "books"
        p = books / "_pre_title_fix_2026-10-17" / "ada_demand_response_kit.pdf"
        if not p.exists():
            self.skipTest("fixture absent")
        r = fte.test_pdf(p)
        self.assertGreaterEqual(r.detail["text_overlaps"], 1)
        self.assertTrue(any("text overlap" in x for x in fte.blocking_problems(r)))
        fixed = books / "ada_demand_response_kit.pdf"
        if fixed.exists():
            self.assertEqual(fte.test_pdf(fixed).detail["text_overlaps"], 0)

    @unittest.skipUnless(importlib.util.find_spec("fitz"),
                         "PyMuPDF not installed: the overlap check is optional in "
                         "file_testing_engine (reports 'PyMuPDF unavailable'); it is "
                         "AGPL-licensed, so adding it to requirements.txt is a "
                         "licensing decision, not a test fix")
    def test_clean_multi_paragraph_pdf_has_no_overlaps(self):
        p = _p("clean_paras.pdf")
        _box_pdf(p, ["First paragraph of ordinary body text."] * 30)
        self.assertEqual(fte.test_pdf(p).detail["text_overlaps"], 0)


class TestPublishGate(unittest.TestCase):
    """distributor.distribute() opens the real file before a real publish."""

    def setUp(self):
        import os
        from channels import registry
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        import test_cycle81 as c81
        self.c81, self.registry, self.os = c81, registry, os
        self._saved = registry.all_arms()
        registry.clear()
        registry.register(c81._fake_arm())
        self.ledger, self.state = str(_p("gate_ledger.jsonl")), str(_p("gate_state.json"))
        for f in (self.ledger, self.state):
            if os.path.exists(f):
                os.remove(f)

    def tearDown(self):
        self.registry.clear()
        for a in self._saved:
            self.registry.register(a)

    def _run(self, path, **kw):
        import distributor
        prod = self.c81._product(title="Dispute Tracker", description="An Excel tracker.",
                                 price_usd=19.0, file_path=str(path))
        # The central Quality Gate is mandatory on real publishes (founder
        # directive, 2026-10-17); this class proves the file-usability gate,
        # so it injects an allowing resolver (see tests/test_gate_mandatory.py).
        kw.setdefault("quality_gate", lambda *a, **k: {
            "allowed": True, "blocked_by": None, "reason": "test allow"})
        return distributor.distribute(prod, arm_names=["gumroad"], allow_protected=True,
                                      ledger_path=self.ledger,
                                      protection_state_path=self.state, **kw)[0]

    def _grid(self):
        p = _p("grid_only.xlsx")
        _skeleton(p, [("Deadlines", ["Dispute ID", "Response due", "Status"], [[PLACEHOLDER]])])
        return p

    def test_formula_less_workbook_blocks_a_real_publish(self):
        out = self._run(self._grid(), dry_run=False)
        self.assertFalse(out["attempted"])
        self.assertIn("file-usability gate", out["skip_reason"])

    def test_dry_run_is_not_gated(self):
        self.assertTrue(self._run(self._grid(), dry_run=True)["attempted"])

    def test_override(self):
        self.assertTrue(self._run(self._grid(), dry_run=False, allow_file_defects=True)["attempted"])

    def test_upgraded_workbook_publishes(self):
        p = self._grid()
        wu.upgrade_workbook(p)
        self.assertTrue(self._run(p, dry_run=False)["attempted"])

    def test_missing_file_does_not_block(self):
        self.assertTrue(self._run(_p("absent.xlsx"), dry_run=False)["attempted"])


if __name__ == "__main__":
    unittest.main()
