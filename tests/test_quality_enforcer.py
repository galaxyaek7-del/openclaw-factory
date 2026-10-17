"""src/quality/quality_enforcer.py: three layers, real files, no network.

Every inspection writes to a temp report dir; payment evidence is passed in
explicitly so no test reads the real market evidence ledger.
"""
import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
spec = importlib.util.spec_from_file_location("quality_enforcer", ROOT / "src" / "quality" / "quality_enforcer.py")
Q = importlib.util.module_from_spec(spec)
spec.loader.exec_module(Q)

GOOD = {
    "id": "T-1", "title": "Freelancer Rate and Raise Calculator", "price": 29,
    "audience": "freelance designers and solo consultants",
    "problem": "they guess their hourly rate and undercharge for years",
    "outcome": "a defensible rate and a dated raise plan in ten minutes",
    "competitors": [{"name": "A", "url": "https://example.com/a", "price_usd": 35},
                    {"name": "B", "url": "https://example.com/b", "price_usd": 49}],
}
PAID = [{"event_type": "closed_sale"}]


def _xlsx(path, formulas=6, sheets=("Inputs", "README")):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheets[0]
    ws.append(["Hours", "Rate", "Total"])
    for i in range(2, 2 + formulas):
        ws.append([10, 50, f"=A{i}*B{i}"])
    for name in sheets[1:]:
        wb.create_sheet(name).append(["How to use: enter your hours in Inputs."])
    wb.save(path)


class TestQualityEnforcer(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.reports = self.tmp / "reports"
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def _inspect(self, product, **kw):
        kw.setdefault("payment_evidence", PAID)
        return Q.inspect(product, report_dir=self.reports, **kw)

    def test_complete_product_is_approved_and_report_is_saved(self):
        x = self.tmp / "calc.xlsx"
        _xlsx(x)
        r = self._inspect(dict(GOOD, files=[str(x)]))
        self.assertEqual(r["verdict"], Q.APPROVED, r["reasons"])
        self.assertFalse(r["blocked"])
        saved = json.loads(Path(r["report_path"]).read_text(encoding="utf-8"))
        self.assertEqual(saved["verdict"], Q.APPROVED)

    def test_no_value_proposition_is_killed(self):
        r = self._inspect({"id": "T-2", "title": "Kit", "audience": "everyone", "price": 0},
                          stage="pre_production")
        self.assertEqual(r["verdict"], Q.KILL)
        self.assertTrue(r["blocked"])

    def test_missing_facts_block_as_missing_data_not_as_pass_or_fail(self):
        p = {k: v for k, v in GOOD.items() if k != "competitors"}
        r = Q.inspect(p, stage="pre_production", payment_evidence=[], report_dir=self.reports)
        self.assertEqual(r["verdict"], Q.BLOCKED_MISSING_DATA)
        self.assertEqual(r["layers"]["market_reality"]["status"], Q.MISSING)
        self.assertIsNone(r["layers"]["market_reality"]["score"])

    def test_thin_text_product_fails_words_per_dollar_rule(self):
        md = self.tmp / "guide.md"
        md.write_text("# Guide\n\nHow to use this guide. " + "Real advice sentence here. " * 40, encoding="utf-8")
        r = self._inspect(dict(GOOD, files=[str(md)]), stage="post_production")
        depth = r["layers"]["content_quality"]["detail"]["depth"]
        self.assertEqual(depth["words_required"], 29 * Q.WORDS_PER_DOLLAR)
        self.assertFalse(depth["text_depth_ok"])
        self.assertEqual(r["layers"]["content_quality"]["score"], 60)   # depth is decisive
        self.assertEqual(r["verdict"], Q.REWRITE)                        # 10 short: fixable

    def test_placeholder_and_broken_file_make_it_not_ready(self):
        md = self.tmp / "draft.md"
        md.write_text("Lorem ipsum TODO [insert intro]", encoding="utf-8")
        bad = self.tmp / "broken.xlsx"
        bad.write_bytes(b"not a workbook")
        r = self._inspect(dict(GOOD, files=[str(md), str(bad)]), stage="post_production")
        cq = r["layers"]["content_quality"]
        self.assertFalse(cq["checks"]["files_open_cleanly"])
        self.assertFalse(cq["checks"]["no_placeholders"])
        self.assertEqual(r["verdict"], Q.NOT_READY)

    def test_unsourced_competitor_prices_are_ignored(self):
        p = dict(GOOD, competitors=[{"name": "no url", "price_usd": 5}, {"name": "no price", "url": "https://x.y"}])
        self.assertEqual(Q.sourced_competitors(p["competitors"]), [])
        r = Q.inspect(p, stage="pre_production", payment_evidence=[], report_dir=self.reports)
        self.assertEqual(r["verdict"], Q.BLOCKED_MISSING_DATA)

    def test_three_sourced_competitors_pass_market_layer_two_do_not(self):
        three = GOOD["competitors"] + [{"name": "C", "url": "https://example.com/c", "price_usd": 39}]
        r3 = Q.inspect(dict(GOOD, competitors=three), stage="pre_production", payment_evidence=[],
                       report_dir=self.reports)
        self.assertEqual(r3["layers"]["market_reality"]["score"], 70)
        self.assertEqual(r3["verdict"], Q.APPROVED)
        r2 = Q.inspect(GOOD, stage="pre_production", payment_evidence=[], report_dir=self.reports)
        self.assertEqual(r2["layers"]["market_reality"]["score"], 60)
        self.assertEqual(r2["verdict"], Q.REWRITE)

    def test_there_is_no_override_argument(self):
        import inspect as pyinspect
        params = set(pyinspect.signature(Q.inspect).parameters)
        self.assertFalse(params & {"override", "force", "bypass", "skip", "advisory"})


if __name__ == "__main__":
    unittest.main()
