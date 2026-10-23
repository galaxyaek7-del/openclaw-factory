"""B13-001 (Renovation Allowance & Change-Order Control) regression guards.

Pins what the first build actually got wrong, so the same class cannot return:

* The first Summary used positional B-references; two of them pointed at the
  wrong row (one at itself), and Excel displayed a blank rather than an error.
  The shipped Summary is therefore re-derived here from typed inputs only.
* The change-order total took the first row instead of the sum, and a value
  labelled as a percentage divided the wrong two numbers.
* The allowance check was a tautology, then loose enough to accept any approved
  change order on the same line. It is now proven by the exact change-order id.
* Loading build_batch12_pdfs.py to reuse its helpers re-emitted five tracked
  PDFs. The guard test pins that the B13-001 loader cannot do it again.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from quality.quality_enforcer import (  # noqa: E402
    _INSTRUCTIONS, _MOJIBAKE, _PLACEHOLDER, _words,
)

XLSX = ROOT / "books" / "batch13" / "B13-001_renovation_allowance_control.xlsx"
PDF = ROOT / "books" / "batch13" / "B13-001_renovation_allowance_control_guide.pdf"
EVID = ROOT / "reports" / "B13-001"

import sys as _gf_sys  # noqa: E402
_gf_sys.path.insert(0, str(Path(__file__).resolve().parent))
from paid_artifacts import paid_files_absent, require_paid_files  # noqa: E402

SHEETS = ["Start Here", "Setup", "Cost Control", "Change Orders", "Allowances",
          "Payments", "Contingency", "Assumptions", "Summary"]

EXPECTED = {
    "original_budget": 45000.0,
    "approved_change_orders": 2300.0,
    "revised_budget": 47300.0,
    "committed_total": 39800.0,
    "paid_to_date": 16300.0,
    "unpaid_committed": 23500.0,
    "forecast_to_complete": 9900.0,
    "projected_final": 49700.0,
    "variance": -2400.0,
    "reserve_remaining": 2400.0,
    "allowance_overruns": 2100.0,
    "credits_owed": 500.0,
    "cos_pending": 1,
    "lines_over": 5,
    "all_in_plan": 49700.0,
    "contingency_drawn": 2100.0,
}


def pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def rederive(wb) -> dict:
    """Independent arithmetic from typed inputs only - never a formula cell."""
    cc, co = wb["Cost Control"], wb["Change Orders"]
    al, pay = wb["Allowances"], wb["Payments"]
    cont, setup = wb["Contingency"], wb["Setup"]

    lines = []
    for r in range(2, 42):
        if not cc.cell(row=r, column=1).value:
            continue
        lines.append({
            "line": cc.cell(row=r, column=1).value,
            "budget": float(cc.cell(row=r, column=4).value or 0),
            "committed": float(cc.cell(row=r, column=5).value or 0),
            "forecast": float(cc.cell(row=r, column=11).value or 0),
        })

    approved, pending = {}, 0
    for r in range(2, 22):
        if not co.cell(row=r, column=1).value:
            continue
        amount = float(co.cell(row=r, column=5).value or 0)
        if str(co.cell(row=r, column=6).value or "").strip().lower() == "yes":
            line = co.cell(row=r, column=3).value
            approved[line] = approved.get(line, 0.0) + amount
        else:
            pending += 1

    paid = {}
    for r in range(2, 32):
        line = pay.cell(row=r, column=2).value
        if line:
            paid[line] = paid.get(line, 0.0) + float(pay.cell(row=r, column=4).value or 0)

    overruns = credits = 0.0
    overage_without_co = 0
    for r in range(2, 22):
        if not al.cell(row=r, column=1).value:
            continue
        allow, chosen = al.cell(row=r, column=4).value, al.cell(row=r, column=5).value
        if allow is None or chosen is None:
            continue
        diff = float(chosen) - float(allow)
        if diff > 0:
            overruns += diff
            if not (al.cell(row=r, column=9).value or "").strip():
                overage_without_co += 1
        elif diff < 0:
            credits += max(0.0, -diff - float(al.cell(row=r, column=7).value or 0))

    reserve = float(setup.cell(row=5, column=2).value or 0)
    drawn = sum(float(cont.cell(row=r, column=3).value or 0)
                for r in range(2, 22) if cont.cell(row=r, column=1).value)

    original = sum(x["budget"] for x in lines)
    revised = original + sum(approved.values())
    committed = sum(x["committed"] for x in lines) + sum(approved.values())
    unpaid = sum(max(0.0, x["committed"] + approved.get(x["line"], 0) - paid.get(x["line"], 0))
                 for x in lines)
    projected = sum(paid.get(x["line"], 0)
                    + max(0.0, x["committed"] + approved.get(x["line"], 0) - paid.get(x["line"], 0))
                    + x["forecast"] for x in lines)
    lines_over = sum(
        1 for x in lines
        if paid.get(x["line"], 0)
        + max(0.0, x["committed"] + approved.get(x["line"], 0) - paid.get(x["line"], 0))
        + x["forecast"] > x["budget"] + approved.get(x["line"], 0))

    return {
        "original_budget": original,
        "approved_change_orders": sum(approved.values()),
        "revised_budget": revised,
        "committed_total": committed,
        "paid_to_date": sum(paid.values()),
        "unpaid_committed": unpaid,
        "forecast_to_complete": sum(x["forecast"] for x in lines),
        "projected_final": projected,
        "variance": revised - projected,
        "reserve_remaining": reserve - drawn,
        "allowance_overruns": overruns,
        "credits_owed": credits,
        "cos_pending": pending,
        "lines_over": lines_over,
        "all_in_plan": revised + reserve - drawn,
        "contingency_drawn": drawn,
        "overage_without_change_order": overage_without_co,
    }


class WorkbookStructureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        require_paid_files(XLSX, PDF)
        cls.wb = load_workbook(XLSX)

    def test_exact_sheet_set_and_order(self):
        self.assertEqual(self.wb.sheetnames, SHEETS)

    def test_enough_real_calculation(self):
        formulas = [c.value for ws in self.wb.worksheets for row in ws.iter_rows()
                    for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertGreaterEqual(len(formulas), 130)
        self.assertGreaterEqual(sum(len(ws.data_validations.dataValidation)
                                    for ws in self.wb.worksheets), 5)

    def test_every_formula_has_balanced_parentheses(self):
        bad = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        if c.value.count("(") != c.value.count(")"):
                            bad.append(f"{ws.title}!{c.coordinate}")
        self.assertEqual(bad, [], "unbalanced formulas make Excel refuse to open the file")

    def test_no_placeholder_words_in_shipped_text(self):
        text = " ".join(str(c.value) for ws in self.wb.worksheets
                        for row in ws.iter_rows() for c in row if isinstance(c.value, str))
        self.assertEqual(_PLACEHOLDER.findall(text), [])

    def test_paid_and_approved_totals_are_formulas_not_typed(self):
        cc = self.wb["Cost Control"]
        for r in (2, 5, 10, 40):
            paid = str(cc.cell(row=r, column=9).value)
            self.assertIn(f"SUMIF(Payments!$B$2:$B$31,$A{r},Payments!$D$2:$D$31)", paid)
            self.assertIn(f"SUMIF('Change Orders'!$C$2:$C$21,$A{r}", str(cc.cell(row=r, column=6).value))
        summary = self.wb["Summary"]
        self.assertIn("SUM('Change Orders'!$G$2:$G$21)", summary["B5"].value)
        self.assertIn("Contingency!$G$2", summary["B14"].value)

    def test_summary_has_no_self_reference(self):
        summary = self.wb["Summary"]
        for row in summary.iter_rows(min_row=4, max_row=20, min_col=2, max_col=2):
            for cell in row:
                value = str(cell.value or "")
                refs = re.findall(r"\$?B\$?(\d+)", value)
                for ref in refs:
                    self.assertNotEqual(int(ref), cell.row,
                                        f"{cell.coordinate} refers to its own row")

    def test_all_in_label_matches_what_it_actually_adds(self):
        # Regression: the row read "(revised budget + reserve)" while adding only
        # the reserve that is left, so a buyer could read it as revised + the full
        # 4,500 reserve (51,800) and be 2,100 short of what the label promised.
        summary = self.wb["Summary"]
        label = str(summary["A19"].value)
        self.assertIn("reserve left", label)
        self.assertNotIn("+ reserve)", label)
        self.assertIn("+B14", str(summary["B19"].value))

    def test_drawn_reserve_is_visible_in_the_summary(self):
        # The 2,100 already drawn from the reserve is spent money that no
        # projected total contains; it must not vanish from the one-page view.
        summary = self.wb["Summary"]
        self.assertEqual(summary["A20"].value, "Contingency drawn to date")
        self.assertEqual(summary["B20"].value, "=Contingency!$H$2")


class ArithmeticTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        require_paid_files(XLSX, PDF)
        cls.model = rederive(load_workbook(XLSX))

    def test_worked_example_rederives_independently(self):
        for key, want in EXPECTED.items():
            self.assertAlmostEqual(self.model[key], want, places=2, msg=key)

    def test_committed_share(self):
        self.assertAlmostEqual(self.model["committed_total"] / self.model["revised_budget"],
                               0.8414, places=3)

    def test_overrun_equals_remaining_reserve_in_the_example(self):
        # Deliberate teaching point: the example's overrun is exactly the reserve
        # that is left, so the reserve is already spoken for.
        self.assertAlmostEqual(-self.model["variance"], self.model["reserve_remaining"], places=2)

    def test_one_allowance_overage_has_no_change_order(self):
        self.assertEqual(self.model["overage_without_change_order"], 1)
        self.assertEqual(self.model["allowance_overruns"], 2100.0)

    def test_allowance_status_is_proven_by_the_change_order_id(self):
        ws = load_workbook(XLSX)["Allowances"]
        self.assertEqual(ws["I2"].value, "CO-01")
        self.assertIn("COUNTIFS('Change Orders'!$A$2:$A$21,$I2",
                      ws["J2"].value)
        self.assertIn('IF($F3>0,IF($I3="","Overage - no change order #"',
                      ws["J3"].value)
        self.assertIn("Credit still owed to you", ws["J3"].value)


class GuideTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        require_paid_files(XLSX, PDF)
        cls.text = pdf_text(PDF)
        cls.flat = re.sub(r"\s+", " ", cls.text)

    def test_depth_meets_the_19_usd_requirement(self):
        self.assertGreaterEqual(len(_words(self.text)), 1900)

    def test_no_placeholders_or_mojibake(self):
        self.assertEqual(_PLACEHOLDER.findall(self.text), [])
        self.assertEqual(_MOJIBAKE.findall(self.text), [])

    def test_it_is_a_guide_with_instructions(self):
        self.assertTrue(_INSTRUCTIONS.search(self.text))
        self.assertGreaterEqual(len(re.findall(r"(?m)^\d{1,2}\. ", self.text)), 12)

    def test_it_names_the_sheets_it_tells_you_to_use(self):
        lowered = self.flat.lower()
        for sheet in ("Cost Control", "Payments", "Change Orders", "Allowances",
                      "Contingency", "Summary", "Assumptions"):
            self.assertIn(sheet.lower(), lowered)

    def test_worked_example_figures_match_the_workbook(self):
        for figure in ("45,000", "2,300", "47,300", "39,800", "16,300", "23,500",
                       "9,900", "49,700", "2,400", "84.1"):
            self.assertIn(figure, self.flat, figure)

    def test_it_states_its_own_limits(self):
        self.assertIn("not legal, tax or financial advice", self.flat)


class ImportGuardTests(unittest.TestCase):
    """The process incident: importing the batch-12 PDF module rebuilt the
    catalogue. The loader must not be able to do that again."""

    def _batch_hashes(self):
        return {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                for p in sorted((ROOT / "books" / "batch12").glob("*.pdf"))}

    def test_guard_flag_stops_the_module_level_builds(self):
        import importlib.util
        import os
        before = self._batch_hashes()
        os.environ["GF_PDF_IMPORT_ONLY"] = "1"
        spec = importlib.util.spec_from_file_location(
            "b12_pdf_guard_probe", ROOT / "scripts" / "build_batch12_pdfs.py")
        module = importlib.util.module_from_spec(spec)
        with self.assertRaises(SystemExit):
            spec.loader.exec_module(module)
        self.assertEqual(before, self._batch_hashes())
        self.assertTrue(hasattr(module, "build"))

    def test_b13_guide_loader_sets_the_flag(self):
        source = (ROOT / "scripts" / "build_b13_001_guide.py").read_text(encoding="utf-8")
        self.assertIn('os.environ["GF_PDF_IMPORT_ONLY"] = "1"', source)
        self.assertIn("except SystemExit", source)


class EvidenceTests(unittest.TestCase):
    def test_native_spec_hash_matches_the_shipped_workbook(self):
        require_paid_files(XLSX)
        import native_workbook_qa as wbqa
        spec = wbqa.SPECS["B13-001"]
        shipped = hashlib.sha256(XLSX.read_bytes()).hexdigest()
        self.assertEqual(spec["artifact_sha256"], shipped)

    def test_gate_evidence_is_a_pass_and_discloses_the_page_policy(self):
        gate = json.loads(next(EVID.glob("GATE_FINAL_*.json")).read_text(encoding="utf-8"))
        self.assertEqual(gate["verdict"], "APPROVED")
        self.assertEqual(gate["price_checked_usd"], 19.0)
        self.assertFalse(gate["page_policy_disclosure"]["market_realism_check_gumroad_digital"]["realistic"])
        self.assertTrue(gate["page_policy_disclosure"]["market_realism_check_gumroad_premium"]["realistic"])

    def test_value_evidence_has_no_mismatch(self):
        value = json.loads(next(EVID.glob("VALUE_INDEPENDENT_CHECK_*.json")).read_text(encoding="utf-8"))
        self.assertEqual(value["verdict"], "PASS")
        self.assertEqual(value["mismatches"], [])
        self.assertTrue(value["self_checks_expected"])

    def test_self_found_defects_are_recorded_including_the_incident(self):
        defects = json.loads(next(EVID.glob("SELF_FOUND_DEFECTS_*.json")).read_text(encoding="utf-8"))
        ids = {d["id"] for d in defects["found_in_own_draft"]}
        self.assertTrue({"D1", "D2", "D3", "D4", "D5"} <= ids)
        self.assertTrue(defects["process_incident"]["disclosed"])


if __name__ == "__main__":
    unittest.main()
