"""spreadsheet_systems produces a genuine .xlsx, never a PDF (2026-10-22).

No network, no LLM. Outputs go to a temp dir.

    python -m pytest tests/test_spreadsheet_family_xlsx.py -q
"""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import asset_generation  # noqa: E402,F401  (registers builders)
import product_packaging.bundle  # noqa: E402,F401
import product_families.families  # noqa: E402,F401
from asset_generation.builders.xlsx_builder import is_genuine_xlsx  # noqa: E402
from product_families import registry  # noqa: E402
from product_families.manifest import get as get_manifest  # noqa: E402

SPEC = {"sheets": [
    {"name": "Cost planner", "columns": ["Item", "Qty", "Unit price", "Line total"],
     "example_rows": [["Fee", 2, 10, ""], ["Label", 3, 4, ""]],
     "formulas": [{"cell": "D2", "formula": "=B2*C2"},
                  {"cell": "Summary!B2", "formula": "=SUM('Cost planner'!D2:D50)", "note": "Total cost"}]},
    {"name": "Log", "columns": ["Date", "Event"], "example_rows": [["2026-10-01", "Started"]]},
]}


class _Tmp(unittest.TestCase):
    def setUp(self):
        self._d = tempfile.TemporaryDirectory()
        self.addCleanup(self._d.cleanup)
        self.d = Path(self._d.name)
        self.fam = registry.get("spreadsheet_systems")


class TestSpreadsheetFamily(_Tmp):
    def test_family_is_registered_and_routes_to_the_xlsx_builder(self):
        self.assertIsNotNone(self.fam)
        m = get_manifest("spreadsheet_systems")
        self.assertEqual(m.asset_builder, "xlsx_workbook")
        self.assertNotEqual(m.asset_builder, "techdoc_package")

    def test_builds_a_genuine_workbook_with_working_formulas(self):
        out = self.d / "planner.xlsx"
        r = self.fam.generate({"workbook_spec": SPEC, "output": str(out)})
        self.assertTrue(r["success"], r.get("error"))
        self.assertTrue(is_genuine_xlsx(out))
        import openpyxl
        wb = openpyxl.load_workbook(out)
        self.assertIn("Cost planner", wb.sheetnames)
        self.assertIn("Summary", wb.sheetnames)
        self.assertEqual(wb["Cost planner"]["D2"].value, "=B2*C2")
        self.assertEqual(wb["Cost planner"]["D3"].value, "=B3*C3")  # row formula filled down
        self.assertEqual(r["qa"]["status"], "PASS")
        self.assertTrue(r["formula_recalculation"].startswith("UNVERIFIED"))
        self.assertEqual(r["package"]["file_count"], 1)

    def test_no_spec_or_no_formulas_fails_closed(self):
        self.assertFalse(self.fam.generate({"output": str(self.d / "a.xlsx")})["success"])
        flat = {"sheets": [{"name": "Grid", "columns": ["A", "B"], "example_rows": [[1, 2]]}]}
        r = self.fam.generate({"workbook_spec": flat, "output": str(self.d / "b.xlsx")})
        self.assertFalse(r["success"])
        self.assertIn("no formulas", r["error"])

    def test_output_must_be_xlsx(self):
        r = self.fam.generate({"workbook_spec": SPEC, "output": str(self.d / "planner.pdf")})
        self.assertFalse(r["success"])


class TestPdfCannotPassAsExcel(_Tmp):
    def _pdf_named_xlsx(self):
        from reportlab.pdfgen import canvas
        p = self.d / "fake.xlsx"
        c = canvas.Canvas(str(p))
        c.drawString(40, 750, "Sheet Structure & Formulas")
        c.save()
        return p

    def test_renamed_pdf_is_not_a_genuine_xlsx(self):
        self.assertFalse(is_genuine_xlsx(self._pdf_named_xlsx()))

    def test_release_gate_rejects_a_pdf_disguised_as_a_workbook(self):
        import release_gate as rg
        r = rg.evaluate("p", [self._pdf_named_xlsx()], "XLSX workbook.", 39, "demand",
                        {"reviewer": "t", "basis": "t"})
        self.assertEqual(r["status"], "QA_FAILED")


if __name__ == "__main__":
    unittest.main()
