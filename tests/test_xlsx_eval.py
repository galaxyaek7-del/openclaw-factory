"""lib/xlsx_eval.py -- independent formula evaluator (2026-10-22).

Expected values follow Excel's documented behavior. Fixtures are written
with openpyxl to a temp dir.

    python -m pytest tests/test_xlsx_eval.py -q
"""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from lib import xlsx_eval as xe  # noqa: E402


def book(cells, extra=None):
    d = tempfile.mkdtemp()
    p = Path(d) / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for coord, v in cells.items():
        ws[coord] = v
    for name, more in (extra or {}).items():
        s = wb.create_sheet(name)
        for coord, v in more.items():
            s[coord] = v
    wb.save(p)
    return xe.Workbook(p)


class TestSemantics(unittest.TestCase):
    def test_arithmetic_and_precedence(self):
        w = book({"A1": 2, "A2": 3, "B1": "=A1+A2*4", "B2": "=-A1^2" if False else "=(A1+A2)*4", "B3": "=10/4"})
        self.assertEqual(w.evaluate("Data!B1"), 14)
        self.assertEqual(w.evaluate("Data!B2"), 20)
        self.assertEqual(w.evaluate("Data!B3"), 2.5)

    def test_blank_cell_is_zero_in_math_and_equal_to_empty_text(self):
        w = book({"B1": "=A1*5", "B2": '=IF(A1="","blank","not")', "B3": "=IF(A1=0,1,0)"})
        self.assertEqual(w.evaluate("Data!B1"), 0)
        self.assertEqual(w.evaluate("Data!B2"), "blank")
        self.assertEqual(w.evaluate("Data!B3"), 1)

    def test_text_in_arithmetic_is_value_error_and_propagates(self):
        w = book({"A1": "x", "B1": "=A1*2", "C1": '=IF(B1>1,"a","b")', "D1": "=IFERROR(B1,0)"})
        self.assertEqual(w.evaluate("Data!B1"), xe.VALUE)
        self.assertEqual(w.evaluate("Data!C1"), xe.VALUE)
        self.assertEqual(w.evaluate("Data!D1"), 0)

    def test_division_by_zero(self):
        self.assertEqual(book({"A1": 0, "B1": "=5/A1"}).evaluate("Data!B1"), xe.DIV0)

    def test_if_is_lazy(self):
        w = book({"A1": 0, "B1": '=IF(A1=0,"safe",1/A1)'})
        self.assertEqual(w.evaluate("Data!B1"), "safe")

    def test_text_vs_number_comparison_follows_excel(self):
        w = book({"A1": "abc", "B1": "=A1>5", "B2": '=A1="ABC"'})
        self.assertIs(w.evaluate("Data!B1"), True)   # text sorts after numbers
        self.assertIs(w.evaluate("Data!B2"), True)   # case-insensitive

    def test_sum_skips_text_and_blanks_in_ranges(self):
        w = book({"A1": 1, "A2": "note", "A4": 5, "B1": "=SUM(A1:A5)"})
        self.assertEqual(w.evaluate("Data!B1"), 6)

    def test_countif_and_cross_sheet_ranges(self):
        w = book({"A1": "kill", "A2": "watch", "A3": "Kill"},
                 {"Summary": {"B1": "=COUNTIF(Data!A1:A9,\"kill\")", "B2": "=COUNTIF('Data'!A1:A9,\"watch\")"}})
        self.assertEqual(w.evaluate("Summary!B1"), 2)
        self.assertEqual(w.evaluate("Summary!B2"), 1)

    def test_and_or_not(self):
        w = book({"A1": 1, "B1": "=AND(A1>0,A1<2)", "B2": "=OR(A1>5,A1=1)", "B3": "=NOT(A1=1)"})
        self.assertIs(w.evaluate("Data!B1"), True)
        self.assertIs(w.evaluate("Data!B2"), True)
        self.assertIs(w.evaluate("Data!B3"), False)

    def test_dates_are_serials(self):
        import datetime
        w = book({"A1": datetime.date(2026, 8, 1), "A2": datetime.date(2026, 10, 15),
                  "B1": "=A1+30", "B2": "=A2-A1"})
        self.assertEqual(w.evaluate("Data!B1"), xe.excel_serial(datetime.date(2026, 8, 31)))
        self.assertEqual(w.evaluate("Data!B2"), 75)
        self.assertEqual(xe.excel_serial(datetime.date(2026, 1, 1)), 46023)  # Excel's own serial

    def test_unsupported_function_is_reported_not_guessed(self):
        with self.assertRaises(xe.Unsupported):
            book({"B1": "=VLOOKUP(1,A1:A2,1,FALSE)"}).evaluate("Data!B1")

    def test_circular_reference_is_reported(self):
        with self.assertRaises(xe.Unsupported):
            book({"A1": "=B1+1", "B1": "=A1+1"}).evaluate("Data!A1")


class TestBuildToolRejectsWrongExpectations(unittest.TestCase):
    def test_a_wrong_expected_value_fails_the_build(self):
        import json
        sys.path.insert(0, str(ROOT / "scripts"))
        import build_tool_workbook as btw
        spec = {"workbook_filename": "t.xlsx", "sheets": [
            {"name": "S", "columns": ["A", "B", "Total"], "example_rows": [[2, 3, ""]],
             "formulas": [{"cell": "C2", "formula": "=A2*B2"}]}],
            "checks": [{"cell": "S!C2", "equals": 7}]}
        d = Path(tempfile.mkdtemp())
        (d / "spec.json").write_text(json.dumps(spec), encoding="utf-8")
        r = btw.build(d / "spec.json", d / "t.xlsx")
        self.assertFalse(r["ok"])
        self.assertEqual(r["recalc"]["failed"][0]["got"], 6)


if __name__ == "__main__":
    unittest.main()
