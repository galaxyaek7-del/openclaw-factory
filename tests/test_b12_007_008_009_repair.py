"""Cycle 106 regression guards for the B12-007 / B12-008 / B12-009 repairs.

Each test pins something that was actually wrong, not a generic assertion:

* B12-008's Summary said "Deposits at risk" over a plain SUM of every deposit,
  so the label described a calculation the formula did not perform. It read
  correctly only because every example row happened to have no photos yet.
* B12-008's guide told the reader to write numbers on the "Dashboard" sheet, but
  that workbook has no Dashboard - the sheet is Summary. A buyer would follow a
  broken instruction.
* B12-008's guide printed a worked mover example that contradicted its own
  scoring sheet (calling a 2,100 quote the maximum score when a 1,950 quote
  scores highest). The corrected arithmetic is now proved against Excel in
  reports/B12-008/WORKED_EXAMPLE_RECONCILIATION_*.json.
* B12-007 and B12-009 were 2-page, ~300-word guides for products priced as full
  guides. Depth is pinned here by the shipped PDF, not by the builder source.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402

import build_batch12_xlsx as b12  # noqa: E402
import native_workbook_qa as wbqa  # noqa: E402

BATCH = ROOT / "books" / "batch12"


def pdf_text(name: str) -> str:
    reader = PdfReader(str(BATCH / name))
    return re.sub(r"\s+", " ", "\n".join((p.extract_text() or "") for p in reader.pages))


def word_count(text: str) -> int:
    """The same tokenizer src/quality/quality_enforcer.py scores depth with."""
    return len(re.findall(r"[^\W_]+", text))


def build_008() -> Path:
    out = Path(tempfile.mkdtemp(prefix="b12_008_test_"))
    return b12.b12_008_moving_budget(out_dir=out)


class MovingBudgetWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.path = build_008()
        cls.wb = load_workbook(cls.path)
        cls.summary = cls.wb["Summary"]

    def _summary_row(self, label: str) -> int:
        for row in self.summary.iter_rows(min_row=2, max_col=3):
            if (row[0].value or "").strip() == label:
                return row[0].row
        self.fail("summary row not found: %r" % label)

    def test_build_writes_outside_the_release_folder(self):
        self.assertNotEqual(self.path.parent, BATCH)

    def test_importing_the_builder_does_not_touch_the_release_artifact(self):
        import importlib
        before = hashlib.sha256((BATCH / "B12-008_moving_budget_tracker.xlsx").read_bytes()).hexdigest()
        importlib.import_module("build_batch12_xlsx")
        after = hashlib.sha256((BATCH / "B12-008_moving_budget_tracker.xlsx").read_bytes()).hexdigest()
        self.assertEqual(before, after)

    def test_deposits_at_risk_is_derived_not_the_gross_total(self):
        row = self._summary_row("Deposits still at risk (USD)")
        self.assertEqual(self.summary.cell(row=row, column=2).value, "=$B$18-$B$19")
        labels = {(r[0].value or "").strip() for r in self.summary.iter_rows(min_row=2, max_col=1)}
        self.assertNotIn("Deposits at risk (USD)", labels)

    def test_every_deposit_metric_has_a_row(self):
        labels = {(r[0].value or "").strip() for r in self.summary.iter_rows(min_row=2, max_col=1)}
        for label in ("Total deposits held (USD)", "Deposits with photos done (USD)",
                      "Deposits still at risk (USD)"):
            self.assertIn(label, labels)

    def test_budget_variance_is_guarded_and_relative_to_the_cheaper_quote(self):
        formula = self.wb["Budget"]["G2"].value
        self.assertIn('IF($E2=""', formula)
        self.assertIn("ISNUMBER($C2)", formula)
        self.assertIn("ISNUMBER($D2)", formula)
        self.assertIn("MIN($C2,$D2)", formula)

    def test_utility_gap_formula_refuses_to_guess(self):
        formula = self.wb["Utilities"]["F2"].value
        self.assertIn("ISNUMBER($D2)", formula)
        self.assertIn("ISNUMBER($E2)", formula)
        self.assertIn('"?"', formula)

    def test_utility_examples_demonstrate_a_gap_and_a_same_day_handover(self):
        ws = self.wb["Utilities"]
        # Two example rows with real dates: the gap formula must have something
        # to compute, otherwise the flagship calculation ships unexercised.
        dated = [r for r in range(2, 5) if ws.cell(row=r, column=4).value and ws.cell(row=r, column=5).value]
        self.assertEqual(len(dated), 2)
        for r in dated:
            self.assertTrue(ws.cell(row=r, column=4).number_format.startswith("yyyy"))
            self.assertTrue(ws.cell(row=r, column=5).number_format.startswith("yyyy"))

    def test_mover_score_is_the_sum_of_three_disclosed_weights(self):
        ws = self.wb["MoverCompare"]
        self.assertEqual(ws["J2"].value, '=IF($B2="","",$G2+$H2+$I2)')
        self.assertIn("MAX(0", ws["G2"].value)
        self.assertIn("MIN($B$2:$B$11)<=0", ws["G2"].value)

    def test_only_the_top_score_is_shortlisted(self):
        ws = self.wb["MoverCompare"]
        for r in range(2, 12):
            self.assertEqual(ws.cell(row=r, column=11).value,
                             '=IF($J{r}="","",IF($J{r}=MAX($J$2:$J$11),"Shortlist",""))'.format(r=r))

    def test_workbook_has_the_depth_the_price_requires(self):
        sheets = self.wb.sheetnames
        self.assertGreaterEqual(len(sheets), 7)
        formulas = sum(1 for ws in self.wb.worksheets for row in ws.iter_rows()
                       for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))
        self.assertGreaterEqual(formulas, 100)
        validations = sum(len(ws.data_validations.dataValidation) for ws in self.wb.worksheets)
        self.assertGreaterEqual(validations, 9)

    def test_no_number_format_is_unterminated(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    fmt = cell.number_format or ""
                    self.assertEqual(fmt.count('"') % 2, 0, "%s!%s has %r" % (ws.title, cell.coordinate, fmt))


class MovingGuideTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.text = pdf_text("B12-008_moving_checklist.pdf")
        cls.lower = cls.text.lower()

    def test_guide_is_deep_enough_for_its_price(self):
        self.assertGreaterEqual(word_count(self.text), 2900)

    def test_guide_points_at_sheets_the_workbook_actually_has(self):
        wb = load_workbook(BATCH / "B12-008_moving_budget_tracker.xlsx", read_only=True)
        try:
            sheets = set(wb.sheetnames)
        finally:
            wb.close()
        for sheet in ("MoverCompare", "AddressChanges", "Utilities", "Summary"):
            self.assertIn(sheet, sheets)
            self.assertIn(sheet.lower(), self.lower, "guide never names the %s sheet" % sheet)
        self.assertNotIn("dashboard", self.lower,
                         "the guide must not name a Dashboard sheet - this workbook has none")

    def test_worked_example_matches_the_scoring_sheet(self):
        # The corrected numbers: cheapest 40+0+30=70, 1,950 insured 34+30+30=94,
        # 2,100 insured 31+30+30=91.
        for needle in ("scores price 34", "94, the top score", "31 + 30 + 30 = 91",
                       "40 x (1 - 250 / 1,700)"):
            self.assertIn(needle, self.text)
        self.assertNotIn("a maximum score", self.text)

    def test_guide_keeps_its_scope_disclaimer(self):
        self.assertIn("NOT LEGAL, FINANCIAL, INSURANCE OR TAX ADVICE", self.text)
        self.assertIn("illustrative arithmetic, not a forecast", self.text)

    def test_guide_never_names_itself_as_a_legal_or_mover_audit(self):
        self.assertIn("does not audit a mover's insurance", self.text)


class DeclutterGuideTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.text = pdf_text("B12-007_digital_declutter_plan.pdf")

    def test_guide_is_deep_enough_for_its_price(self):
        self.assertGreaterEqual(word_count(self.text), 1900)

    def test_guide_promises_only_the_three_verdicts(self):
        self.assertIn("exactly one verdict", self.text)
        self.assertIn("Archive is the pressure valve", self.text)

    def test_guide_names_only_sheets_the_workbook_has(self):
        wb = load_workbook(BATCH / "B12-007_digital_inventory.xlsx", read_only=True)
        try:
            sheets = set(wb.sheetnames)
        finally:
            wb.close()
        for sheet in sheets:
            if sheet == "README":
                continue
            self.assertIn(sheet.lower(), self.text.lower(),
                          "guide never mentions the %s sheet it depends on" % sheet)

    def test_backup_section_separates_sync_from_backup(self):
        self.assertIn("cloud sync folder is not a backup", self.text)
        self.assertIn("Perform a real restore test", self.text)


class HandoverKitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.text = pdf_text("B12-009_pet_sitter_handover_kit.pdf")

    def test_kit_is_deep_enough_for_its_price(self):
        self.assertGreaterEqual(word_count(self.text), 1900)

    def test_emergency_protocol_precedes_contacting_the_owner(self):
        self.assertIn("Call the vet before you call the owner", self.text)
        for step in range(1, 7):
            self.assertIn("Step %d." % step, self.text)

    def test_kit_covers_the_forms_a_sitter_actually_fills_in(self):
        for needle in ("Spending limit I authorise", "Who may authorise treatment",
                       "Daily log", "Medication given"):
            self.assertIn(needle, self.text)

    def test_kit_refuses_to_let_a_sitter_improvise_treatment(self):
        self.assertIn("Never give human medication", self.text)
        self.assertIn("Never induce vomiting unless a vet", self.text)
        self.assertIn("not to double the next one", self.text)

    def test_prose_errors_fixed_this_cycle_stay_fixed(self):
        self.assertIn("an eye that cannot open", self.text)
        self.assertNotIn("a eye that cannot open", self.text)
        self.assertIn("section eleven", self.text)
        self.assertNotIn("section six records", self.text)

    def test_kit_states_its_own_limits(self):
        self.assertIn("NOT VETERINARY OR LEGAL ADVICE", self.text)
        self.assertIn("instructions always override anything written here", self.text)


class EvidenceConsistencyTests(unittest.TestCase):
    """The recorded evidence must describe the artifact that would ship."""

    def test_registered_specs_point_at_real_files_with_matching_hashes(self):
        for product_id, spec in wbqa.SPECS.items():
            artifact = ROOT / spec["artifact"]
            self.assertTrue(artifact.exists(), product_id)
            self.assertEqual(hashlib.sha256(artifact.read_bytes()).hexdigest(),
                             spec["artifact_sha256"], product_id)

    def test_native_qa_evidence_matches_the_shipping_workbook(self):
        artifact = ROOT / wbqa.B12_008_SPEC["artifact"]
        digest = hashlib.sha256(artifact.read_bytes()).hexdigest()
        records = sorted((ROOT / "reports" / "B12-008").glob("WORKBOOK_QA_*.json"))
        self.assertTrue(records, "no native workbook QA evidence recorded")
        record = json.loads(records[-1].read_text(encoding="utf-8"))
        self.assertEqual(record["verdict"], "PASS")
        self.assertEqual(record["artifact_sha256"], digest)
        self.assertEqual(record["critical_count"], 0)


if __name__ == "__main__":
    unittest.main()
