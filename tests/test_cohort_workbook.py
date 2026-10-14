"""Tests for cohort_workbook.build_xlsx -- one per defect it fixes.

Each defect was found in shipped workbooks on 2026-10-14; the specs below are
reduced copies of the real ones (GF-P10-10, GF-B11-06).
"""
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import cohort_workbook as cw

TMP = Path(tempfile.gettempdir()) / "gf_cohort_workbook"


def _build(spec):
    TMP.mkdir(parents=True, exist_ok=True)
    sp, xp = TMP / "spec.json", TMP / "out.xlsx"
    sp.write_text(json.dumps(spec), encoding="utf-8")
    rep = cw.build_xlsx(sp, xp)
    return openpyxl.load_workbook(xp), rep


PIPELINE = {"sheets": [{
    "name": "Deal_Pipeline",
    "columns": ["Deal_ID", "Brand", "Stage", "Fee", "Next_Action_Date"],
    "example_rows": [["D-201", "Northwind Outdoors", "Negotiating", 2500, "2026-09-28"],
                     ["D-202", "Lumen Skincare", "Paid", 1800, "2026-10-05"]],
    "formulas": [{"cell": "Summary!B2", "formula": "=COUNTA('Deal_Pipeline'!A2:A200)",
                  "note": "Total deals tracked"}],
    "notes": "Stage values: Outreach, Paid."}]}


class TestSummaryAddressedFormula(unittest.TestCase):
    def test_brand_is_not_overwritten_and_summary_exists(self):
        wb, rep = _build(PIPELINE)
        self.assertEqual(wb["Deal_Pipeline"]["B2"].value, "Northwind Outdoors")
        self.assertIn("Summary", wb.sheetnames)
        s = wb["Summary"]
        self.assertEqual(s["A2"].value, "Total deals tracked")
        self.assertEqual(s["B2"].value, "=COUNTA('Deal_Pipeline'!A2:A200)")

    def test_notes_go_to_a_notes_sheet_not_under_the_table(self):
        wb, _ = _build(PIPELINE)
        self.assertEqual(wb["Notes"]["B1"].value, "Stage values: Outreach, Paid.")
        self.assertEqual(wb["Deal_Pipeline"].max_row, 3)

    def test_iso_dates_become_real_dates(self):
        wb, rep = _build(PIPELINE)
        v = wb["Deal_Pipeline"]["E2"].value
        self.assertEqual(v.date() if isinstance(v, dt.datetime) else v, dt.date(2026, 9, 28))
        self.assertEqual(rep["dates"], 2)


TRACKER = {"sheets": [{
    "name": "Deal_Tracker",
    "columns": ["Deal ID", "Accepted", "Closing", "Days"],
    "example_rows": [["D-01", "2026-09-02", "2026-10-17", 45],
                     ["D-02", "2026-09-05", "2026-10-20", 40]],   # 40 is wrong (should be 45)
    "formulas": [{"cell": "D2", "formula": "=C2-B2"},
                 {"cell": "A8", "formula": "=COUNTA(A2:A20)"}]}]}


class TestInTableAggregateAndFillDown(unittest.TestCase):
    def test_aggregate_in_a_data_column_moves_to_summary_qualified(self):
        wb, rep = _build(TRACKER)
        self.assertIsNone(wb["Deal_Tracker"]["A8"].value)
        self.assertIn("Deal_Tracker!A8", rep["moved"])
        self.assertEqual(wb["Summary"]["B2"].value, "=COUNTA('Deal_Tracker'!A2:A20)")

    def test_row_formula_fills_entry_rows(self):
        wb, rep = _build(TRACKER)
        ws = wb["Deal_Tracker"]
        self.assertEqual(ws["D2"].value, "=C2-B2")
        self.assertEqual(ws["D4"].value, "=C4-B4")
        self.assertEqual(ws["D50"].value, "=C50-B50")

    def test_authored_example_values_are_never_overwritten(self):
        wb, rep = _build(TRACKER)
        self.assertEqual(wb["Deal_Tracker"]["D3"].value, 40)
        self.assertEqual(rep["mismatches"][0]["cell"], "Deal_Tracker!D3")


class TestFormulaHelpers(unittest.TestCase):
    def test_qualify_leaves_string_literals_alone(self):
        f = '=COUNTIF(F2:F30,"A1 MISSING")'
        self.assertEqual(cw.qualify(f, "Doc"), "=COUNTIF('Doc'!F2:F30,\"A1 MISSING\")")

    def test_row_formula_detection(self):
        self.assertTrue(cw.is_row_formula("=B2*E2", 2))
        self.assertTrue(cw.is_row_formula('=IF(H2="","",H2-G2)', 2))
        self.assertFalse(cw.is_row_formula("=SUM(B2:B9)", 2))
        # a same-row range is a row total (Risk_Scorer H2 in GF-P10-10), not a summary
        self.assertTrue(cw.is_row_formula("=SUM(B2:G2)", 2))
        self.assertFalse(cw.is_row_formula("=B2*'Rate'!B2", 2))
        self.assertFalse(cw.is_row_formula("=B3*E2", 2))

    def test_shift_row(self):
        self.assertEqual(cw.shift_row('=IF(F2<>"","OK",C2-TODAY())', 2, 7),
                         '=IF(F7<>"","OK",C7-TODAY())')


if __name__ == "__main__":
    unittest.main()
