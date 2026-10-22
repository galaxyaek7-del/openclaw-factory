#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Galaxy Forge -- build and verify a working .xlsx tool from a spec (2026-10-22).

    python scripts/build_tool_workbook.py books/_workbook_specs/<name>.json [--out books/<name>.xlsx]

Steps, failing closed at each one:
  1. build: product_families' spreadsheet_systems (asset_generation's
     xlsx_workbook -> cohort_workbook.build_xlsx), which itself requires
     formulas and a passing file_testing_engine check;
  2. post: number formats and list validations from spec["post"];
  3. recalc: every spec["checks"] cell is recalculated by lib/xlsx_eval.py (an
     independent evaluator; no native engine is available here) and compared with the value the
     spec author computed by hand. Any mismatch fails the build.

Specs live in books/_workbook_specs/ (gitignored: they are the paid
product). Prints one JSON report; exit 1 on any failure.
"""
import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

TOL = 1e-4


def _apply_post(path, post):
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.load_workbook(path)
    for sheet, cols in (post.get("number_formats") or {}).items():
        ws = wb[sheet]
        for col, fmt in cols.items():
            for row in range(2, ws.max_row + 1):
                ws[f"{col}{row}"].number_format = fmt
    for sheet, cols in (post.get("validations") or {}).items():
        ws = wb[sheet]
        for col, options in cols.items():
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            dv.add(f"{col}2:{col}{max(ws.max_row, 201)}")
            ws.add_data_validation(dv)
    wb.save(path)


def _recalc(path, checks):
    """Recalculate with lib/xlsx_eval.py, a small independent evaluator (no
    native engine is available here: no LibreOffice; pycel fails on Python
    3.14; `formulas` needs a scipy DLL blocked by application control). An
    unsupported construct fails the check as UNVERIFIED, never as a pass."""
    from lib import xlsx_eval
    wb = xlsx_eval.Workbook(path)
    results = []
    for c in checks:
        try:
            got = wb.evaluate(c["cell"])
        except xlsx_eval.Unsupported as e:
            results.append({"cell": c["cell"], "ok": False, "error": f"UNVERIFIED: {e}"})
            continue
        want = c["equals"]
        if isinstance(want, (int, float)) and isinstance(got, (int, float)) and not isinstance(got, bool):
            ok = abs(float(got) - float(want)) <= max(TOL, abs(want) * 1e-6)
        else:
            ok = str(got) == str(want)
        results.append({"cell": c["cell"], "ok": ok, "expected": want, "got": repr(got) if not isinstance(got, (int, float, str)) else got})
    return results


def build(spec_path, out=None):
    import asset_generation  # noqa: F401  (registers builders)
    import product_families.families  # noqa: F401
    import product_packaging.bundle  # noqa: F401
    from product_families import registry

    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    out = Path(out) if out else ROOT / "books" / spec["workbook_filename"]
    report = {"spec": str(spec_path), "output": str(out), "ok": False}
    built = registry.get("spreadsheet_systems").generate({"workbook_spec": spec, "output": str(out)})
    report["build"] = {k: built.get(k) for k in ("success", "error")}
    report["file_test"] = (built.get("qa") or {}).get("status")
    if not built.get("success"):
        return report
    _apply_post(out, spec.get("post") or {})
    import file_testing_engine as fte
    report["file_test_after_post"] = fte.test_xlsx(out).status
    checks = _recalc(out, spec.get("checks") or [])
    report["recalc"] = {"engine": "lib/xlsx_eval.py", "checked": len(checks),
                        "failed": [c for c in checks if not c["ok"]]}
    report["ok"] = (report["file_test_after_post"] == "PASS" and bool(checks)
                    and not report["recalc"]["failed"])
    if not checks:
        report["recalc"]["note"] = "no checks declared: recalculation UNVERIFIED"
    return report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--out")
    a = ap.parse_args()
    r = build(a.spec, a.out)
    print(json.dumps(r, ensure_ascii=False, default=str, indent=1))
    sys.exit(0 if r["ok"] else 1)


if __name__ == "__main__":
    main()
