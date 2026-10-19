#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Galaxy Forge Batch 12 — XLSX builders (products 001/002-reg/004/005-tracker/
006-log/007-inv/008-budget/010-pipe). Deterministic, $0, no network.
Writes into books/batch12/. Each workbook carries real structure + working
formulas where math is the product; example rows are labeled EXAMPLE.
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "books" / "batch12"
OUT.mkdir(parents=True, exist_ok=True)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
EX_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Border(*(Side(style="thin", color="B0B0B0") for _ in range(4)))
WRAP = Alignment(wrap_text=True, vertical="top")


def _paint_formulas(wb):
    """Paint every formula cell light blue so README color language is
    true in every workbook: white = fill me, blue = computed (do not
    overwrite), yellow = EXAMPLE row (overwrite or delete). Applied last
    so it wins over example-row fills on formula cells."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.fill = FORMULA_FILL


def _sheet(wb, title, headers, rows, widths=None, formatters=None, example_rows=0):
    ws = wb.active if wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 else wb.create_sheet(title)
    if ws.title != title:
        ws.title = title
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = HDR_FONT, HDR_FILL, WRAP
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border, cell.alignment = THIN, WRAP
            if formatters and c - 1 in formatters:
                cell.number_format = formatters[c - 1]
    for r in range(2, 2 + example_rows):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = EX_FILL
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr = None
    return ws


def _readme(wb, lines):
    ws = wb.create_sheet("README")
    ws["A1"] = lines[0]
    ws["A1"].font = Font(bold=True, size=13)
    for i, ln in enumerate(lines[1:], 3):
        ws.cell(row=i, column=1, value=ln).alignment = WRAP
    ws.column_dimensions["A"].width = 110



def _b12_guard_inputs(ws):
    """Cycle 94 (edge cases found in Excel 16): typing 30 instead of 30% gave
    $2.82/hr silently, tax >= 100% or negative inputs gave negative rates, and
    0 weeks/hours or text gave #DIV/0! / #VALUE!. Validate the inputs and make
    every output say what to fix instead of showing an error or a wrong number."""
    from openpyxl.worksheet.datavalidation import DataValidation
    rules = [
        ("C2", "decimal", 0, 10_000_000, "Take-home pay", "Enter a yearly amount in USD, 0 or more."),
        ("C3", "decimal", 0, 0.9, "Tax buffer", "Enter a percentage between 0% and 90%, e.g. 30% (not 30)."),
        ("C4", "decimal", 0, 10_000_000, "Expenses", "Enter yearly expenses in USD, 0 or more."),
        ("C5", "whole", 1, 52, "Billable weeks", "Enter a whole number of weeks between 1 and 52."),
        ("C6", "decimal", 1, 80, "Billable hours", "Enter billable hours per week between 1 and 80."),
        ("C7", "decimal", 0, 1, "Safety margin", "Enter a percentage between 0% and 100%, e.g. 15% (not 15)."),
    ]
    for cell, kind, lo, hi, title, msg in rules:
        dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi),
                            allow_blank=False, showErrorMessage=True, showInputMessage=True,
                            errorTitle=title, error=msg, promptTitle=title, prompt=msg)
        ws.add_data_validation(dv)
        dv.add(cell)
    ws["D8"] = ('=IF(AND(ISNUMBER(C2),ISNUMBER(C3),ISNUMBER(C4),C2>=0,C4>=0,C3>=0,C3<1),'
                'C2/(1-C3)+C4,"Check take-home, tax % (0-90%) and expenses")')
    ws["D9"] = ('=IF(NOT(ISNUMBER(D8)),"Fix the line above",'
                'IF(AND(ISNUMBER(C5),ISNUMBER(C6),C5>0,C6>0),D8/(C5*C6),"Check weeks and hours"))')
    ws["D10"] = ('=IF(NOT(ISNUMBER(D9)),"Fix the lines above",'
                 'IF(AND(ISNUMBER(C7),C7>=0),D9*(1+C7),"Check safety margin %"))')
    # A tax buffer entered as 30 instead of 30% is the most likely mistake.
    ws["E3"] = '=IF(AND(ISNUMBER(C3),C3>0.9),"Tax buffer must be 90% or less. Type 30% (not 30).","")'

def b12_001_rate_calculator():
    """Freelancer Rate & Raise Calculator: cost-plus rate math + utilization."""
    wb = Workbook()
    _sheet(wb, "Inputs", ["#", "Input", "Your value", "Notes"], [
        [1, "Target annual take-home pay (USD)", 80000, "After tax, what you want to keep"],
        [2, "Tax + self-employment buffer %", 0.30, "EXAMPLE 30% — set your own rate"],
        [3, "Annual business expenses (USD)", 6000, "Tools, insurance, accounting"],
        [4, "Billable weeks per year", 46, "52 minus vacation, gaps, admin"],
        [5, "Billable hours per week", 25, "Realistic delivery hours, not desk hours"],
        [6, "Utilization safety margin %", 0.15, "Buffer for scope creep and late payers"],
    ], [6, 42, 16, 60], {2: '#,##0'}, example_rows=0)
    ws = wb["Inputs"]
    ws["C8"] = "Gross revenue needed (USD)"
    ws["D8"] = "=C2/(1-C3)+C4"
    ws["C8"].font = Font(bold=True)
    ws["C9"] = "Minimum hourly rate (USD/hr)"
    ws["D9"] = "=D8/(C5*C6)"
    ws["C9"].font = Font(bold=True)
    ws["D9"].number_format = '"$"#,##0.00'
    ws["C10"] = "Rate with safety margin (USD/hr)"
    ws["D10"] = "=D9*(1+C7)"
    ws["D10"].number_format = '"$"#,##0.00'
    ws["C10"].font = Font(bold=True)
    # Cycle 88 fix: the column format '#,##0' made 30% and 15% display as "0".
    for c in ("C3", "C7"):
        ws[c].number_format = "0%"
    ws["D8"].number_format = '"$"#,##0.00'
    _b12_guard_inputs(ws)
    _sheet(wb, "ClientMix", ["Client", "Hours/mo", "Rate/hr (USD)", "Monthly revenue (USD)", "Hassle score 1-5", "Effective rank"],
        [["EXAMPLE Acme Co", 40, 95, "=B2*C2", 2, "=D2/E2"],
         ["EXAMPLE Beta LLC", 20, 120, "=B3*C3", 4, "=D3/E3"]],
        [22, 12, 16, 22, 16, 16], {1: '#,##0', 2: '"$"#,##0', 3: '"$"#,##0.00', 5: '"$"#,##0.00'}, example_rows=2)
    _sheet(wb, "RaisePlanner", ["Trigger to review", "Metric to check", "Raise %", "New rate (USD/hr)", "Date done"],
        # Cycle 88 fix: new rate follows the user's own rate (Inputs!D10), not a hardcoded 95.
        # Cycle 94: shows a message instead of an error when Inputs is incomplete.
        [["EXAMPLE quarterly review", "Utilization > 85% for 6 weeks", 0.10,
          '=IF(ISNUMBER(Inputs!$D$10),Inputs!$D$10*(1+C2),"Fix Inputs first")', ""],
         ["EXAMPLE new skill ships", "Client-paid premium work", 0.15,
          '=IF(ISNUMBER(Inputs!$D$10),Inputs!$D$10*(1+C3),"Fix Inputs first")', ""]],
        [26, 34, 10, 18, 14], {2: "0%", 3: '"$"#,##0.00'}, example_rows=2)
    _readme(wb, ["B12-001 Freelancer Rate & Raise Calculator — HOW TO USE",
        "1) Fill ONLY the white input cells (column C on Inputs). Blue cells are computed — do not overwrite them. Yellow rows are EXAMPLES — overwrite with your numbers or delete them.",
        "2) Your charge rate = Inputs!D10. Quote at or above it.",
        "3) ClientMix: log every client monthly; sort by Effective rank (revenue per hassle point). Fire the bottom row first.",
        "4) RaisePlanner: review on the triggers listed; never raise without a logged trigger.",
        "This is arithmetic on YOUR inputs, not financial advice. Tax buffer is an estimate — confirm with your accountant."])
    p = OUT / "B12-001_freelancer_rate_raise_calculator.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_002_handoff_register():
    """Client Project Handoff Register (Cycle 98 rebuild to the deep-QA bar).

    Register + per-project asset checklist + risk register + a computed close
    gate + a dashboard of transfer readiness. White cells are yours to fill,
    blue cells are formulas (do not overwrite), yellow rows are EXAMPLES.
    Every rollup guards its inputs, so an empty register shows a blank or a
    plain-language prompt instead of an Excel error."""
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()

    # ---- HandoffRegister ---------------------------------------------
    reg_headers = ["#", "Project", "Outgoing owner", "Incoming owner", "Handoff date",
                   "Days to handoff", "Status", "Assets complete %", "Open high risks",
                   "Client sign-off (Y/N)", "Close gate"]
    reg_rows = [
        [1, "EXAMPLE Website rebuild", "You", "Client team", "2026-11-01", None, "In progress", None, None, "N", None],
        [2, "EXAMPLE API integration", "You", "Agency B", "2026-11-15", None, "Not started", None, None, "N", None],
    ] + [[i, "", "", "", None, None, "", None, None, "", None] for i in range(3, 16)]
    _sheet(wb, "HandoffRegister", reg_headers, reg_rows,
           [4, 26, 16, 16, 13, 14, 13, 15, 14, 16, 16],
           {4: "yyyy-mm-dd", 5: "0", 7: "0%"}, example_rows=2)
    reg = wb["HandoffRegister"]
    for r in range(2, 16):
        reg[f"F{r}"] = f'=IF(ISNUMBER(E{r}),E{r}-TODAY(),"")'
        reg[f"H{r}"] = (f'=IF($B{r}="","",IFERROR(COUNTIFS(AssetChecklist!$A$2:$A$60,$B{r},'
                        f'AssetChecklist!$D$2:$D$60,"Y")/COUNTIFS(AssetChecklist!$A$2:$A$60,$B{r}),""))')
        reg[f"I{r}"] = (f'=IF($B{r}="","",COUNTIFS(RiskRegister!$A$2:$A$40,$B{r},'
                        f'RiskRegister!$H$2:$H$40,"OPEN",RiskRegister!$F$2:$F$40,">=12"))')
        reg[f"K{r}"] = (f'=IF($B{r}="","",IF($H{r}="","Fill the asset checklist first",'
                        f'IF(AND($H{r}>=1,$I{r}=0,$J{r}="Y"),"READY TO CLOSE","OPEN")))')
    dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,1,1)",
                        allow_blank=True, showErrorMessage=True, errorTitle="Handoff date",
                        error="Enter a real date, e.g. 2026-11-01.")
    reg.add_data_validation(dv); dv.add("E2:E15")
    dv = DataValidation(type="list", formula1='"Not started,In progress,Ready to close,Closed"',
                        allow_blank=True, showErrorMessage=True, errorTitle="Status",
                        error="Choose one of: Not started, In progress, Ready to close, Closed.")
    reg.add_data_validation(dv); dv.add("G2:G15")
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Client sign-off", error="Enter Y or N.")
    reg.add_data_validation(dv); dv.add("J2:J15")

    # ---- ChecklistTemplate (the canonical items) ---------------------
    _sheet(wb, "ChecklistTemplate", ["Category", "Item to transfer", "Why it prevents a dispute"],
        [["Code", "Repository ownership moved to the client org; their admins, not yours", "You keep write access and the 'transferred' claim is false"],
         ["Code", "Deploy keys / CI secrets reissued to the client and the old ones revoked", "A live key in your hands is a latent security incident"],
         ["Code", "Staging and production reproduced by the incoming owner, live, on the call", "A system they cannot run is not a system they own"],
         ["Docs", "Architecture notes: where everything lives and what depends on what", "The next integrator starts from zero without this map"],
         ["Docs", "Runbook: deploy, roll back, restart, rotate credentials", "An incident with no runbook becomes your unpaid emergency"],
         ["Docs", "Known-bugs list, written honestly, with severity", "Undisclosed bugs become 'you hid this' arguments"],
         ["Accounts", "Hosting, DNS, analytics, error tracking: transfer ownership, never share a password", "Shared passwords blur who changed what"],
         ["Accounts", "Every third-party API key reissued to a client-owned account", "The client must not depend on your account staying open"],
         ["Accounts", "Billing for the above moved to the client, verified on the call", "A silent recurring charge on your card is a dispute in month two"],
         ["Money", "Final invoice issued and payment terms stated in writing", "Scope disputes are usually money disputes in disguise"],
         ["Money", "Outstanding expenses and retainer balance reconciled to the cent", "Open money is the most common reason a handoff stalls"],
         ["Data", "Client data exported in an open format and delivered, not left in your tool", "Data held hostage is the worst handoff failure"],
         ["Data", "Your working copies of client data deleted, confirmed and dated", "A copy you forgot about is a privacy liability"],
         ["Security", "Every shared login replaced with a named account for each person", "Anonymous shared logins can never be revoked or audited"],
         ["Security", "Access you still hold revoked or explicitly time-boxed", "Standing access after handoff is a finding in any audit"],
         ["People", "Stakeholder introductions done with one line of context each", "Names without context are useless to the receiver"],
         ["People", "One named decision-maker on each side, in writing", "Two owners means no owner"],
         ["Support", "Support window: dates, what is covered, response time, and what costs extra", "An undefined window becomes free, permanent, unlimited support"]],
        [12, 56, 46], example_rows=0)
    _sheet(wb, "AssetChecklist", ["Project", "Category", "Item", "Access transferred (Y/N)",
                                  "Verified by", "Verified date", "Where it lives (URL / path)", "Proof note"],
        [["EXAMPLE Website rebuild", "Code", "Repo ownership + deploy keys", "N", "", None, "github.com/…", ""],
         ["EXAMPLE Website rebuild", "Accounts", "DNS registrar ownership", "N", "", None, "registrar login", ""],
         ["EXAMPLE Website rebuild", "Money", "Final invoice issued", "N", "", None, "", ""]] +
        [["", "", "", "", "", None, "", ""] for _ in range(34)],
        [24, 12, 46, 22, 14, 13, 30, 26], {5: "yyyy-mm-dd"}, example_rows=3)
    ac = wb["AssetChecklist"]
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Access transferred", error="Enter Y or N.")
    ac.add_data_validation(dv); dv.add("D2:D60")
    dv = DataValidation(type="list",
                        formula1='"Code,Docs,Accounts,Money,Data,Security,People,Support"',
                        allow_blank=True, showErrorMessage=True, errorTitle="Category",
                        error="Choose a category: Code, Docs, Accounts, Money, Data, Security, People, Support.")
    ac.add_data_validation(dv); dv.add("B2:B60")

    # ---- RiskRegister -------------------------------------------------
    _sheet(wb, "RiskRegister", ["Project", "Risk (what could go wrong)", "Owner", "Severity 1-5",
                                "Likelihood 1-5", "Score (S×L)", "Mitigation", "Status (OPEN/CLOSED)", "Due date"],
        [["EXAMPLE Website rebuild", "DNS transfer fails during the window", "You", 4, 3, None,
          "Lower TTL 48h before; keep old host warm for 72h", "OPEN", "2026-10-30"],
         ["", "", "", None, None, None, "", "", None]] +
        [["", "", "", None, None, None, "", "", None] for _ in range(24)],
        [22, 40, 12, 12, 14, 12, 40, 20, 12], {3: "0", 4: "0", 8: "yyyy-mm-dd"}, example_rows=1)
    rr = wb["RiskRegister"]
    for r in range(2, 40):
        rr[f"F{r}"] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}*E{r},"")'
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=True, showErrorMessage=True, errorTitle="Severity",
                        error="Enter a whole number from 1 (low) to 5 (severe).")
    rr.add_data_validation(dv); dv.add("D2:D40")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=True, showErrorMessage=True, errorTitle="Likelihood",
                        error="Enter a whole number from 1 (rare) to 5 (almost certain).")
    rr.add_data_validation(dv); dv.add("E2:E40")
    dv = DataValidation(type="list", formula1='"OPEN,CLOSED"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Status", error="Enter OPEN or CLOSED.")
    rr.add_data_validation(dv); dv.add("H2:H40")
    rr["J1"] = "A score of 12 or more (S×L) counts as a HIGH risk and blocks the close gate until it is CLOSED."
    rr["J1"].font = Font(italic=True)

    # ---- Dashboard ----------------------------------------------------
    _sheet(wb, "Dashboard", ["KPI", "Value", "How it is computed"],
        [["Projects tracked", '=COUNTA(HandoffRegister!$B$2:$B$15)', "Named rows on HandoffRegister"],
         ["Projects READY TO CLOSE", '=COUNTIF(HandoffRegister!$K$2:$K$15,"READY TO CLOSE")',
          "Sign-off Y, assets 100%, no open high risks"],
         ["Sign-offs collected", '=COUNTIF(HandoffRegister!$J$2:$J$15,"Y")', "Client sign-off = Y"],
         ["Open HIGH risks", '=COUNTIFS(RiskRegister!$F$2:$F$40,">=12",RiskRegister!$H$2:$H$40,"OPEN")',
          "Score >= 12 and status OPEN"],
         ["Assets transferred (Y)", '=COUNTIF(AssetChecklist!$D$2:$D$60,"Y")', "Access transferred = Y"],
         ["Assets outstanding (N or blank)", '=COUNTA(AssetChecklist!$C$2:$C$60)-COUNTIF(AssetChecklist!$D$2:$D$60,"Y")',
          "Named checklist items not yet transferred"],
         ["Transfer readiness %", '=IFERROR(COUNTIF(AssetChecklist!$D$2:$D$60,"Y")/COUNTA(AssetChecklist!$C$2:$C$60),"")',
          "Transferred ÷ named items"],
         ["Next handoff in (days)", '=IFERROR(MIN(HandoffRegister!$F$2:$F$15),"")',
          "Smallest non-empty Days to handoff"]],
        [34, 18, 52], {1: "0"}, example_rows=0)
    wb["Dashboard"]["B8"].number_format = "0%"

    _readme(wb, ["B12-002 Client Project Handoff Register — HOW TO USE",
        "1) HandoffRegister: one row per project. Fill only the white cells. Blue cells compute (Days to handoff, Assets complete %, Open high risks, Close gate) and must not be overwritten. Yellow rows are EXAMPLES — overwrite or delete them.",
        "2) AssetChecklist: add one row per item you are transferring, and set its Project name so it rolls up. Copy the exact items you need from ChecklistTemplate (Code, Docs, Accounts, Money, Data, Security, People, Support). Nothing counts as transferred until Access transferred = Y.",
        "3) RiskRegister: log every handoff risk, score it Severity x Likelihood, and CLOSE the ones scoring 12 or more before you close the project. The close gate stays OPEN while a high risk is open.",
        "4) Dashboard: reads the three sheets and shows transfer readiness, sign-offs, and the next handoff date. A project is 'READY TO CLOSE' only when its assets are 100%, it has no open high risks, and the client has signed off.",
        "5) Use with the B12-002 PDF guide (the 45-minute session script and the copy-paste emails). This is a record-keeping tool, not legal advice."])
    p = OUT / "B12-002_client_handoff_register.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_004_meeting_calculator():
    """Meeting Cost & ROI Calculator (Cycle 97 rebuild to the deep-QA bar).

    Inputs -> outputs are separated: white cells are yours to fill, blue
    cells are computed, yellow rows are EXAMPLES. Every output guards its
    inputs so a mis-typed percentage or a blank row shows a message
    instead of a wrong number or an Excel error."""
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()

    # ---- Setup --------------------------------------------------------
    _sheet(wb, "Setup", ["#", "Setting", "Your value", "Notes"], [
        [1, "Default loaded hourly cost (USD/hr)", 65,
         "Fully loaded = salary / 2080 x (1 + surcharge). A team median is fine; use the helper below for exact."],
        [2, "Working weeks per year", 46, "52 minus vacation, holidays and non-billable gaps"],
        [3, "Currency label", "USD", "Change the label only; the numbers stay plain"],
        [4, "Decision value threshold (USD per decision)", 500,
         "Above this cost per decision, the meeting is flagged Shorten/async"],
        [5, "Optional: annual salary (USD)", 100000, "Used only to derive a loaded rate below"],
        [6, "Optional: employer surcharge % (EU ~20%, US ~35%)", 0.30, "Taxes, benefits and overhead on top of salary"],
        [7, "Loaded $/hr (computed from row 5-6)", "=C6/2080*(1+C7)",
         "Copy this number into MeetingLog column C, or just type your own rate there"],
    ], [5, 46, 16, 70], None, example_rows=0)
    setup = wb["Setup"]
    # Per-cell formats (the value column mixes a rate, a count, a label and %).
    for cell in ("C2", "C5", "C6", "C8"):
        setup[cell].number_format = '"$"#,##0.00'
    setup["C3"].number_format = '0'
    setup["C7"].number_format = '0%'
    for cell, kind, lo, hi, title, msg in [
        ("C2", "decimal", 0, 100000, "Loaded hourly rate",
         "Enter the fully-loaded cost per hour in USD (0 or more), e.g. 65."),
        ("C3", "whole", 1, 52, "Working weeks",
         "Enter a whole number of weeks between 1 and 52."),
        ("C5", "decimal", 0, 10000000, "Decision threshold",
         "Enter a USD amount, 0 or more."),
        ("C6", "decimal", 0, 10000000, "Annual salary",
         "Enter a yearly salary in USD, 0 or more."),
        ("C7", "decimal", 0, 1, "Surcharge %",
         "Enter a percentage between 0% and 100%, e.g. 30% (not 30)."),
    ]:
        dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi),
                            allow_blank=True, showErrorMessage=True, showInputMessage=True,
                            errorTitle=title, error=msg, promptTitle=title, prompt=msg)
        setup.add_data_validation(dv)
        dv.add(cell)

    # ---- MeetingLog ---------------------------------------------------
    headers = ["Meeting", "Attendees", "Loaded $/hr", "Hours/meeting", "Times/week",
               "Decisions/meeting", "Weekly cost (USD)", "Annual cost (USD)",
               "Person-hours/year", "Cost per decision (USD)", "Suggested action"]
    rows = []
    for r in range(2, 22):
        if r == 2:
            rows.append(["EXAMPLE Sprint planning", 6, 75, 1.5, 1, 2, None, None, None, None, None])
        elif r == 3:
            rows.append(["EXAMPLE Status round-robin", 8, 70, 0.5, 3, 0, None, None, None, None, None])
        else:
            rows.append([""] * 11)
    _sheet(wb, "MeetingLog", headers, rows,
           [26, 11, 13, 14, 12, 18, 17, 17, 16, 20, 22],
           {2: '"$"#,##0.00', 3: "0.0", 4: "0.0", 6: '"$"#,##0.00', 7: '"$"#,##0.00',
            8: '#,##0', 9: '"$"#,##0.00'}, example_rows=2)
    log = wb["MeetingLog"]
    for r in range(2, 22):
        log[f"G{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),ISNUMBER(D{r}),ISNUMBER(E{r})),B{r}*C{r}*D{r}*E{r},"")'
        log[f"H{r}"] = f'=IF(ISNUMBER(G{r}),G{r}*Setup!$C$3,"")'
        log[f"I{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(D{r}),ISNUMBER(E{r})),B{r}*D{r}*E{r}*Setup!$C$3,"")'
        log[f"J{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),ISNUMBER(D{r}),ISNUMBER(F{r}),F{r}>0),B{r}*C{r}*D{r}/F{r},"")'
        log[f"K{r}"] = (f'=IF(NOT(ISNUMBER(G{r})),"",'
                        f'IF(AND(ISNUMBER(F{r}),F{r}=0),"KILL or make async",'
                        f'IF(AND(ISNUMBER(J{r}),J{r}>Setup!$C$5),"Shorten or async","Keep")))')
    for cell, kind, lo, hi, title, msg in [
        ("B2:B21", "whole", 1, 500, "Attendees", "Enter a whole number of attendees, 1 or more."),
        ("C2:C21", "decimal", 0, 10000, "Loaded $/hr", "Enter the fully-loaded hourly cost in USD, 0 or more."),
        ("D2:D21", "decimal", 0.05, 24, "Hours", "Enter hours per meeting (0.05 to 24), e.g. 0.5 for 30 min."),
        ("E2:E21", "decimal", 0, 40, "Times/week", "Enter how many times this meeting runs per week (0 to 40)."),
        ("F2:F21", "whole", 0, 100, "Decisions/meeting", "Enter how many real decisions this meeting produces (0 if none)."),
    ]:
        dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi),
                            allow_blank=True, showErrorMessage=True, showInputMessage=True,
                            errorTitle=title, error=msg, promptTitle=title, prompt=msg)
        log.add_data_validation(dv)
        dv.add(cell)

    # ---- Dashboard ----------------------------------------------------
    _sheet(wb, "Dashboard", ["KPI", "Value", "Note"], [
        ["Meetings logged", '=COUNTA(MeetingLog!A2:A21)', "Rows with a meeting name on MeetingLog"],
        ["Total weekly meeting cost (USD)", '=IFERROR(SUM(MeetingLog!G2:G21),0)', "Sum of the Weekly cost column"],
        ["Monthly meeting cost (USD)", '=IFERROR(B3*4.33,0)', "Weekly x 4.33 weeks"],
        ["Annual meeting cost (USD)", '=IFERROR(SUM(MeetingLog!H2:H21),0)', "Weekly x working weeks per year (Setup)"],
        ["Total person-hours per year", '=IFERROR(SUM(MeetingLog!I2:I21),0)', "Attendees x hours x times/week x weeks"],
        ["Zero-decision meetings", '=SUMPRODUCT((MeetingLog!A2:A21<>"")*(MeetingLog!F2:F21=0))',
         "Logged meetings that produce no decision - your first cuts"],
        ["Annual cost of zero-decision meetings (USD)",
         '=IFERROR(SUMIFS(MeetingLog!H2:H21,MeetingLog!F2:F21,0,MeetingLog!A2:A21,"<>"),0)',
         "The saving if you cancel every zero-decision meeting"],
        ["Most expensive meeting", '=IFERROR(INDEX(MeetingLog!A2:A21,MATCH(MAX(MeetingLog!H2:H21),MeetingLog!H2:H21,0)),"")',
         "Highest annual cost"],
        ["Highest annual meeting cost (USD)", '=IFERROR(MAX(MeetingLog!H2:H21),0)', ""],
        ["Highest cost per decision (USD)", '=IFERROR(MAX(MeetingLog!J2:J21),"")',
         "The meeting doing the least work per dollar"],
    ], [40, 20, 52], {2: '"$"#,##0.00'}, example_rows=0)
    dash = wb["Dashboard"]
    for r in (3, 4, 5, 8, 10, 11):
        dash[f"B{r}"].number_format = '"$"#,##0.00'

    # ---- Rules --------------------------------------------------------
    _sheet(wb, "Rules", ["#", "Rule", "Applies to"],
        [[1, "No agenda in invite = decline by default", "All recurring"],
         [2, "Default 25 min, not 30; 50 min, not 60", "All new invites"],
         [3, "Status updates go async (doc/chat), never a meeting", "Status round-robins"],
         [4, "Owner + decision + deadline in every invite", "Decision meetings"],
         [5, "Quarterly kill review: cancel the costliest meetings with zero decisions", "Recurring"],
         [6, "Set an end date on every recurring meeting and revisit it", "Recurring"]],
        [5, 62, 20], example_rows=0)

    _readme(wb, ["B12-004 Meeting Cost & ROI Calculator — HOW TO USE",
        "1) Setup: set your default loaded hourly rate and working weeks. Loaded cost = salary / 2080 x (1 + surcharge). The helper derives it from a salary if you prefer.",
        "2) MeetingLog: one row per recurring meeting. Fill ONLY the white cells (attendees, rate, hours, times/week, decisions/meeting). Blue cells compute cost and are not to be overwritten. Yellow rows are EXAMPLES - overwrite or delete them.",
        "3) Cost per decision = one meeting's cost divided by the real decisions it produces. It is the number worth arguing about.",
        "4) Dashboard: totals, the most expensive meeting, and the annual saving from cancelling every zero-decision meeting. Start your kill review at the top of the Meetings logged list.",
        "5) Rules: the operating defaults. Re-run the quarterly kill review or the cost quietly returns.",
        "Arithmetic on YOUR inputs only. Loaded rates and decisions are estimates you own; this is not HR, legal or financial advice."])
    p = OUT / "B12-004_meeting_cost_roi_calculator.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_005_guest_tracker():
    wb = Workbook()
    _sheet(wb, "Pipeline", ["Show", "Contact", "Pitch date", "Angle", "Stage", "Record date", "Follow-up date", "Assets sent (Y/N)"],
        [["EXAMPLE The Ops Show", "producer@example.com", "2026-10-20", "Freelancer pricing math", "Pitched", "", "2026-10-27", "N"]],
        [22, 26, 12, 30, 14, 13, 15, 16], example_rows=1)
    _sheet(wb, "Angles", ["#", "Your angle (one line)", "Proof you can cite", "Best-fit show type"],
        [[1, "EXAMPLE pricing math freelancers get wrong", "500 rate reviews analyzed", "Freelance/business"],
         [2, "", "", ""],
         [3, "", "", ""]],
        [5, 44, 34, 24], example_rows=1)
    _sheet(wb, "FollowUps", ["Show", "Touch #", "Date", "Channel", "One-line value add", "Reply (Y/N)"],
        [["EXAMPLE The Ops Show", 1, "2026-10-27", "Email", "EXAMPLE rate-benchmark table", "N"]],
        [22, 9, 12, 10, 44, 11], example_rows=1)
    _readme(wb, ["B12-005 Podcast Guest Booking Tracker — HOW TO USE",
        "1) Write 3+ angles first (Angles sheet) — never pitch without an angle.",
        "2) Log every pitch on Pipeline; move FollowUps on day 7 and day 21, then close.",
        "3) Pair with the B12-005 PDF (pitch + follow-up scripts word-for-word). Max 2 follow-ups per show — more is spam."])
    p = OUT / "B12-005_podcast_guest_tracker.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_006_home_office_log():
    wb = Workbook()
    _sheet(wb, "Expenses2026", ["Date", "Category", "Description", "Amount (USD)", "Home %", "Deductible est. (USD)", "Receipt file"],
        [["2026-01-15", "EXAMPLE internet", "Monthly bill", 80.00, 0.40, "=D2*E2", "receipts/jan-internet.pdf"]],
        [12, 18, 30, 15, 10, 19, 28], {3: '"$"#,##0.00', 5: '"$"#,##0.00'}, example_rows=1)
    ws = wb["Expenses2026"]
    ws["H1"] = "TOTAL deductible est. (USD)"
    ws["I1"] = "=SUM(F2:F500)"
    ws["I1"].number_format = '"$"#,##0.00'
    _sheet(wb, "SpaceMeasure", ["Item", "Value", "Notes"],
        [["Total home sq ft", 1200, "Measure once"], ["Office sq ft (exclusive use)", 150, "Must be regular + exclusive use"],
         ["Office % (= office ÷ total)", 0.125, "=B3/B2 — used for Home % above"], ["Method chosen", "Actual", "Actual or Simplified — pick ONE per year"]],
        [30, 14, 44], example_rows=0)
    _sheet(wb, "Categories", ["Category", "What counts", "What never counts"],
        [["Internet/utilities", "Business-use share only", "Full bill without home %"],
         ["Supplies", "100% business items", "Personal items"],
         ["Repairs", "Office-area repairs", "Whole-home improvements (different rules)"],
         ["Rent/mortgage share", "Home % of eligible costs (actual method)", "Simplified method mixes"]],
        [20, 44, 44], example_rows=0)
    _readme(wb, ["B12-006 Home-Office Deduction Organizer — HOW TO USE",
        "1) Measure once (SpaceMeasure). Pick ONE method for the year: Actual (this log) or Simplified.",
        "2) Log every expense with Home % and receipt file. The deductible estimate column is arithmetic, not a filing.",
        "3) Read the companion PDF before filing. PROCEDURE INFORMATION ONLY — NOT TAX ADVICE. Verify current-year IRS rules; limits change."])
    p = OUT / "B12-006_home_office_expense_log.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_007_digital_inventory():
    wb = Workbook()
    for sheet, headers, rows in [
        ("Accounts", ["Service", "Username", "2FA (Y/N)", "Recovery email OK (Y/N)", "Keep / Delete", "Deleted date"],
         [["EXAMPLE old-forum.com", "me2009", "N", "N", "Delete", ""]]),
        ("Files", ["Location", "Size est.", "Last opened", "Keep / Archive / Delete", "Done date"],
         [["EXAMPLE downloads/2019", "4 GB", "2024-02-01", "Archive", ""]]),
        ("Backups", ["What", "Where (3-2-1 slot)", "Last verified restore", "Status"],
         [["EXAMPLE photos", "Cloud + external drive", "never", "Unverified"]]),
        ("Subscriptions", ["Service", "Cost/mo (USD)", "Last used", "Keep / Cancel", "Cancelled date"],
         [["EXAMPLE streaming-x", 15.99, "2026-08-01", "Cancel", ""]]),
    ]:
        _sheet(wb, sheet, headers, rows, [26, 18, 22, 22, 16, 14], example_rows=1)
    _readme(wb, ["B12-007 Digital Inventory — HOW TO USE",
        "1) Work the 30-day plan PDF week by week; every decision lands in one of these four sheets.",
        "2) 3-2-1 backup rule: 3 copies, 2 media, 1 offsite. A backup never restore-tested = no backup.",
        "3) Never store passwords in this file — it is an inventory (what/where/status), not a vault."])
    p = OUT / "B12-007_digital_inventory.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_008_moving_budget():
    wb = Workbook()
    _sheet(wb, "Budget", ["Item", "Quote A (USD)", "Quote B (USD)", "Chosen (USD)", "Paid (Y/N)"],
        [["EXAMPLE movers (3-bed)", 1800, 2100, 1800, "N"],
         ["EXAMPLE packing supplies", 150, 120, 120, "N"],
         ["EXAMPLE deposit + first month", 3000, 3000, 3000, "N"],
         ["EXAMPLE utility transfers/fees", 200, 200, 200, "N"]],
        [32, 16, 16, 15, 11], {1: '"$"#,##0', 2: '"$"#,##0', 3: '"$"#,##0'}, example_rows=4)
    ws = wb["Budget"]
    ws["A8"] = "TOTAL chosen (USD)"
    ws["D8"] = "=SUM(D2:D7)"
    ws["D8"].number_format = '"$"#,##0'
    _sheet(wb, "MoverCompare", ["Mover", "Quote (USD)", "Insured (Y/N)", "Reviews checked (Y/N)", "Red flags", "Decision"],
        [["EXAMPLE MoveRight", 1800, "Y", "Y", "None", "Shortlist"]],
        [20, 14, 13, 20, 30, 12], example_rows=1)
    _sheet(wb, "AddressChanges", ["Account/service", "Updated (Y/N)", "Date", "Confirmation #"],
        [["EXAMPLE bank", "N", "", ""], ["EXAMPLE employer/payroll", "N", "", ""],
         ["EXAMPLE subscriptions", "N", "", ""], ["EXAMPLE voter registration", "N", "", ""]],
        [30, 14, 12, 20], example_rows=4)
    _readme(wb, ["B12-008 Moving Budget & Tracker — HOW TO USE",
        "1) Get 2+ quotes per row (MoverCompare). Never book an uninsured mover.",
        "2) Work AddressChanges in week 2 — missed ones cost money for months.",
        "3) Pair with the B12-008 8-week checklist PDF."])
    p = OUT / "B12-008_moving_budget_tracker.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_010_grant_pipeline():
    wb = Workbook()
    _sheet(wb, "Pipeline", ["Funder", "Program", "Amount (USD)", "Deadline", "Stage", "Owner", "Docs checklist done (Y/N)", "Submitted date"],
        [["EXAMPLE Community Fund", "Small grants", 10000, "2027-02-01", "Research", "Sam", "N", ""]],
        [22, 20, 14, 12, 14, 10, 22, 15], {2: '"$"#,##0'}, example_rows=1)
    _sheet(wb, "Stages", ["Stage", "Exit rule", "Typical docs"],
        [["Research", "Funder fits mission + budget range", "Guidelines PDF saved"],
         ["Eligibility", "All must-pass criteria = yes", "Eligibility checklist"],
         ["Drafting", "Narrative + budget + attachments complete", "Narrative, budget, board list, 990"],
         ["Review", "Second reader signed off", "Review sign-off"],
         ["Submitted", "Confirmation receipt saved", "Receipt"],
         ["FollowUp", "Report dates calendared", "Award letter / decline note"]],
        [14, 40, 44], example_rows=0)
    _sheet(wb, "Boilerplate", ["Block", "Owner", "Last updated", "Status"],
        [["EXAMPLE mission paragraph", "", "", "Missing"],
         ["EXAMPLE budget template", "", "", "Missing"],
         ["EXAMPLE board list", "", "", "Missing"]],
        [28, 12, 14, 12], example_rows=3)
    _readme(wb, ["B12-010 Grant Pipeline Tracker — HOW TO USE",
        "1) Nothing enters Drafting until Eligibility exits clean — ineligible applications are the #1 wasted effort.",
        "2) Reuse Boilerplate blocks; never rewrite mission/budget prose per funder from scratch.",
        "3) Log every decline with the reason — a pipeline without decline reasons cannot improve."])
    p = OUT / "B12-010_grant_pipeline_tracker.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


def b12_003_access_continuity_inventory():
    """Founder Access Continuity Inventory (Cycle 98 build to the deep-QA bar).

    An inventory, never a vault: it records WHAT exists, WHERE it lives and WHO
    may touch it - never a password. White cells are yours to fill, blue cells
    are formulas (do not overwrite), yellow rows are EXAMPLES. Every rollup
    guards its inputs, so an empty sheet shows a blank or a plain-language
    prompt instead of an Excel error."""
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()

    # ---- Systems: what exists, where it lives, who may touch it -------
    _sheet(wb, "Systems", ["#", "What it is", "Where it lives (provider / URL)",
                           "Account owner", "Recovery route (how access is restored)",
                           "Two-person rule (Y/N)", "Criticality 1-5", "Last verified", "Notes"],
        [[1, "EXAMPLE domain registrar", "registrar.example (login: owner@business)",
          "You", "Email reset to a shared inbox two people can open", "N", 5, None,
          "Never store the password here - only the route"],
         [2, "EXAMPLE business bank", "Bank (relationship manager: [name])",
          "You", "Branch visit with ID + account number", "Y", 5, None, ""],
         [3, "EXAMPLE payment processor", "payouts.biz (payouts to the business account)",
          "You", "Support ticket from the recovery email", "Y", 4, None, ""]] +
        [[i, "", "", "", "", "", None, None, ""] for i in range(4, 31)],
        [4, 24, 40, 16, 44, 18, 14, 13, 30], {6: "0", 7: "yyyy-mm-dd"}, example_rows=3)
    sy = wb["Systems"]
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Two-person rule", error="Enter Y or N.")
    sy.add_data_validation(dv); dv.add("F2:F30")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=True, showErrorMessage=True, errorTitle="Criticality",
                        error="Enter a whole number 1 (nice to have) to 5 (business stops without it).")
    sy.add_data_validation(dv); dv.add("G2:G30")
    sy["J1"] = ("A password never goes in this sheet. Record the RECOVERY ROUTE only "
                "(how access is restored), never the secret itself.")
    sy["J1"].font = Font(italic=True)

    # ---- Holders: who may act, and how far -----------------------------
    _sheet(wb, "Holders", ["Role", "Name", "Best contact", "What they may touch (categories)",
                           "Authority basis (named on account / will / POA)", "Confirmed (Y/N)", "Last confirmed"],
        [["EXAMPLE executor / successor", "[Full name]", "phone + email",
          "Bank, registrar, processor", "Named on the account / will", "N", None],
         ["EXAMPLE backup operator", "[Full name]", "phone + email",
          "Hosting, email, file storage", "Named on the account", "N", None],
         ["EXAMPLE accountant", "[Firm + person]", "email",
          "Bookkeeping + tax filings", "Engagement letter", "N", None]] +
        [["", "", "", "", "", "", None] for _ in range(12)],
        [24, 20, 22, 34, 38, 14, 13], {6: "yyyy-mm-dd"}, example_rows=3)
    ho = wb["Holders"]
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Confirmed", error="Enter Y or N.")
    ho.add_data_validation(dv); dv.add("F2:F15")
    ho["H1"] = ("Authority like a power of attorney or a will is a legal matter - "
                "get it drawn up by a lawyer where you live. This sheet records who you told, not legal authority.")
    ho["H1"].font = Font(italic=True)

    # ---- Envelopes: the three-envelope method (doctrine, not a record) --
    _sheet(wb, "Envelopes", ["Envelope", "What it holds", "Where it is kept", "Who holds a copy",
                             "How it is opened (process, never a secret)", "Last checked"],
        [["A - the map", "This inventory: what exists, where it lives, who may touch it",
          "With this workbook (printed + the file)", "", "Read directly - it contains no secrets", None],
         ["B - the keys", "Password-manager emergency access (the vault, not the passwords)",
          "Password manager's own emergency-access setting", "",
          "The named contact requests access; the manager's waiting period runs; then they set their own password", None],
         ["C - the letter", "Master-password hint (never the password), lawyer + accountant contacts, 30-day keep-alive steps",
          "Sealed, with a lawyer or in a safe", "",
          "Opened only by the named successor, or on a stated condition", None]],
        [16, 52, 34, 18, 52, 13], {5: "yyyy-mm-dd"}, example_rows=0)

    # ---- DrillLog: the annual 30-minute drill --------------------------
    _sheet(wb, "DrillLog", ["Date", "Run by", "A up to date (Y/N)", "B flow re-tested (Y/N)",
                            "C location known (Y/N)", "Contacts current (Y/N)", "Gaps found",
                            "Fixed on", "Next drill due"],
        [[None, "", "", "", "", "", "EXAMPLE the recovery email was a dead address", None, None]] +
        [[None, "", "", "", "", "", "", None, None] for _ in range(9)],
        [13, 14, 15, 17, 18, 18, 44, 13, 14], {0: "yyyy-mm-dd", 7: "yyyy-mm-dd", 8: "yyyy-mm-dd"},
        example_rows=1)
    dl = wb["DrillLog"]
    for r in range(2, 11):
        dl[f"I{r}"] = f'=IF(ISNUMBER(A{r}),EDATE(A{r},12),"")'
    for col in ("C", "D", "E", "F"):
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Drill result", error="Enter Y or N.")
        dl.add_data_validation(dv); dv.add(f"{col}2:{col}10")

    # ---- Dashboard -----------------------------------------------------
    _sheet(wb, "Dashboard", ["KPI", "Value", "How it is computed"],
        [["Systems inventoried", '=COUNTA(Systems!$B$2:$B$30)', "Named rows on Systems"],
         ["Systems under a two-person rule", '=COUNTIF(Systems!$F$2:$F$30,"Y")',
          "Two-person rule = Y"],
         ["Two-person coverage %",
          '=IFERROR(COUNTIF(Systems!$F$2:$F$30,"Y")/COUNTA(Systems!$B$2:$B$30),"")',
          "Covered systems / named systems"],
         ["Critical systems (4-5) with NO two-person rule",
          '=COUNTIFS(Systems!$G$2:$G$30,">=4",Systems!$F$2:$F$30,"<>Y")',
          "The rows that would stop the business and have no backup person"],
         ["Holders named", '=COUNTA(Holders!$B$2:$B$15)', "Named rows on Holders"],
         ["Holders confirmed", '=COUNTIF(Holders!$F$2:$F$15,"Y")', "Confirmed = Y"],
         ["Last drill date", '=IFERROR(MAX(DrillLog!$A$2:$A$10),"")', "Most recent drill"],
         ["Days since last drill", '=IF(ISNUMBER(B8),TODAY()-B8,"")', "Today - last drill date"],
         ["Drill status", '=IF(NOT(ISNUMBER(B8)),"NEVER DRILLED",IF(B9>365,"OVERDUE","Current"))',
          "Overdue once more than 365 days pass"],
         ["Envelopes checked (all 3)", '=SUMPRODUCT((Envelopes!$F$2:$F$4<>"")*1)', "Envelope rows with a checked date"],
         ["Continuity readiness",
          '=IF(B2=0,"Fill the Systems sheet",'
          'IF(AND(B3>=1,ISNUMBER(B9),B9<=365),'
          'IF(B5=0,"CONTINUITY READY","GAPS REMAIN - close critical gaps"),"GAPS REMAIN"))',
          "A drill within 12 months, at least one backup person, no high-criticality gap"]],
        [42, 22, 54], None, example_rows=0)
    dash = wb["Dashboard"]
    dash["B3"].number_format = "0"
    dash["B4"].number_format = "0%"
    dash["B5"].number_format = "0"
    dash["B6"].number_format = "0"
    dash["B7"].number_format = "0"
    dash["B8"].number_format = "yyyy-mm-dd"
    dash["B9"].number_format = "0"
    dash["B11"].number_format = "0"

    _readme(wb, ["B12-003 Founder Access Continuity Inventory - HOW TO USE",
        "1) This is an INVENTORY, never a vault. Never type a password, PIN or full account number into any cell here. You record WHAT exists, WHERE it lives and WHO may touch it - plus the recovery route. The passwords live in a real password manager.",
        "2) Systems: one row per thing the business runs on (registrar, hosting, bank, processor, email, storage). Fill only the white cells. Criticality 4-5 means the business stops without it. Two-person rule = Y when more than one trusted person can restore access. Blue cells on Dashboard compute; do not overwrite them.",
        "3) Holders: who may act, what they may touch, and on what basis. Keep the confirmed column honest - re-confirm every year. Legal authority (will, power of attorney) must be drawn up by a lawyer where you live; this sheet only records who you have told.",
        "4) Envelopes: the three-envelope method. A = this map (no secrets). B = password-manager emergency access (the named contact requests it and the manager runs its waiting period). C = a sealed letter with a hint and the keep-alive steps, held by a lawyer or in a safe.",
        "5) DrillLog: once a year, run the 30-minute drill, record what broke, and fix it. The Dashboard marks the drill OVERDUE after 12 months - an untested plan is a wish, not a plan.",
        "6) Pair this with the B12-003 PDF guide (the method, the successor's first 72 hours, and the fill-in tables). This is an organising system, not legal, financial or insurance advice."])
    p = OUT / "B12-003_access_continuity_inventory.xlsx"
    _paint_formulas(wb)
    wb.save(p)
    return p


if __name__ == "__main__":
    makers = [b12_001_rate_calculator, b12_002_handoff_register, b12_003_access_continuity_inventory,
              b12_004_meeting_calculator,
              b12_005_guest_tracker, b12_006_home_office_log, b12_007_digital_inventory,
              b12_008_moving_budget, b12_010_grant_pipeline]
    for fn in makers:
        p = fn()
        print("built", p.name, p.stat().st_size, "bytes")
